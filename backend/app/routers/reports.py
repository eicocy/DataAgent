from __future__ import annotations

from datetime import UTC, datetime
from urllib.parse import quote
import io

from fastapi import APIRouter, Depends, HTTPException, Query, Request
from fastapi.responses import Response
from pydantic import BaseModel, ConfigDict, Field
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.artifacts.storage import LocalArtifactStorage
from app.config import get_settings
from app.database import get_db
from app.dependencies import current_user
from app.models import (AnalysisArtifact, AnalysisRecord, AnalysisReport,
                        AnalysisReportVersion, ToolExecutionRecord, User)
from app.reports.service import ReportError, create_report, enqueue_report_task
from app.reports.schemas import ReportSpec
from app.artifacts.manager import ArtifactManager, expired


router = APIRouter(tags=["analysis reports and artifacts"])


class ReportRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")
    session_id: int = Field(gt=0)
    source_record_ids: list[int] = Field(min_length=1, max_length=50)
    spec: ReportSpec
    request_id: str | None = Field(default=None, min_length=1, max_length=64)


class ExportRequest(BaseModel):
    model_config = ConfigDict(extra='forbid')
    request_id: str | None = Field(default=None, min_length=1, max_length=64)


def _report(db: Session, report_id: int, user_id: int) -> AnalysisReport:
    report = db.get(AnalysisReport, report_id)
    if not report or report.user_id != user_id:
        raise HTTPException(404, detail={"code": "REPORT_NOT_FOUND", "message": "报告不存在"})
    return report


def _version(db: Session, report_id: int, number: int | None) -> AnalysisReportVersion:
    query = select(AnalysisReportVersion).where(AnalysisReportVersion.report_id == report_id)
    if number is None:
        query = query.order_by(AnalysisReportVersion.version_number.desc())
    else:
        query = query.where(AnalysisReportVersion.version_number == number)
    row = db.scalar(query.limit(1))
    if not row:
        raise HTTPException(404, detail={"code": "REPORT_VERSION_NOT_FOUND", "message": "报告版本不存在"})
    return row


def _error(error: ReportError):
    message = {
        "REPORT_VERSION_CONFLICT": "报告已被其他编辑更新，请刷新后再保存",
        "REPORT_SOURCE_REQUIRED": "请选择至少一条成功的分析记录",
        "REPORT_SOURCE_NOT_FOUND": "报告引用的分析记录不可用",
        "REPORT_SOURCE_MISMATCH": "编辑报告时不能更换会话或数据版本",
    }.get(error.code, "报告请求无法完成")
    raise HTTPException(error.status_code, detail={"code": error.code, "message": message}) from None


def _view(report: AnalysisReport, version: AnalysisReportVersion) -> dict:
    return {"id": report.id, "session_id": report.session_id, "dataset_id": report.dataset_id,
            "dataset_version_id": report.dataset_version_id, "title": report.title,
            "status": report.status, "latest_version": report.latest_version_number,
            "version": version.version_number, "spec": version.spec_json,
            "document": version.document_json, "source_record_ids": version.source_records_json,
            "created_at": version.created_at.isoformat()}


def _owned_artifact(db: Session, artifact_id: int, user_id: int) -> AnalysisArtifact:
    artifact = db.get(AnalysisArtifact, artifact_id)
    if not artifact or artifact.user_id != user_id:
        raise HTTPException(404, detail={"code": "ARTIFACT_NOT_FOUND", "message": "文件不存在"})
    if artifact.report_version_id:
        version = db.get(AnalysisReportVersion, artifact.report_version_id)
        report = db.get(AnalysisReport, version.report_id) if version else None
        if not report or report.user_id != user_id:
            raise HTTPException(404, detail={"code": "ARTIFACT_NOT_FOUND", "message": "文件不存在"})
    elif artifact.record_id:
        record = db.get(AnalysisRecord, artifact.record_id)
        if not record or record.user_id != user_id:
            raise HTTPException(404, detail={"code": "ARTIFACT_NOT_FOUND", "message": "文件不存在"})
    elif artifact.tool_execution_id:
        execution = db.get(ToolExecutionRecord, artifact.tool_execution_id)
        if not execution or execution.user_id != user_id:
            raise HTTPException(404, detail={"code": "ARTIFACT_NOT_FOUND", "message": "文件不存在"})
    else:
        raise HTTPException(404, detail={"code": "ARTIFACT_NOT_FOUND", "message": "文件不存在"})
    return artifact


def _artifact_view(item: AnalysisArtifact) -> dict:
    now = datetime.now(UTC).replace(tzinfo=None)
    is_expired = expired(item)
    return {"artifact_id": item.id, "task_id": item.record_id or item.tool_execution_id or item.report_version_id,
            "conversation_id": None, "artifact_type": item.kind.upper(),
            "status": "EXPIRED" if is_expired else item.status, "title": item.file_name or item.step_id,
            "mime_type": item.mime_type or ("application/json" if item.stored_name.endswith(".json") else "application/octet-stream"),
            "file_name": item.file_name, "size_bytes": item.size_bytes,
            "preview_url": f"/api/v1/artifacts/{item.id}/preview",
            "download_url": f"/api/v1/artifacts/{item.id}/download",
            "metadata": item.metadata_json or {}, "created_at": item.created_at.isoformat(),
            "expires_at": item.expires_at.isoformat() if item.expires_at else None}


@router.post("/reports", status_code=201)
def create(request: ReportRequest, response: Response, user: User = Depends(current_user), db: Session = Depends(get_db)):
    try:
        record = enqueue_report_task(db, user.id, request.session_id, request.spec.dataset_id,
            request.spec.dataset_version_id, "create", {"spec": request.spec.model_dump(mode="json"),
            "source_record_ids": request.source_record_ids}, request_id=request.request_id)
    except ReportError as error:
        _error(error)
    response.status_code = 200 if getattr(record, '_request_replayed', False) else 202
    return {"code": response.status_code, "message": "accepted", "data": {"record_id": record.id,
        "session_id": record.session_id, "status": record.status,
        "status_url": f"/api/v1/analysis/runs/{record.id}",
        "events_url": f"/api/v1/analysis/runs/{record.id}/events"}}


@router.get("/reports")
def list_reports(session_id: int | None = Query(default=None, gt=0),
                 user: User = Depends(current_user), db: Session = Depends(get_db)):
    query = select(AnalysisReport).where(AnalysisReport.user_id == user.id)
    if session_id:
        query = query.where(AnalysisReport.session_id == session_id)
    rows = db.scalars(query.order_by(AnalysisReport.updated_at.desc()).limit(100)).all()
    return {"code": 200, "message": "success", "data": [
        {"id": row.id, "session_id": row.session_id, "dataset_id": row.dataset_id,
         "dataset_version_id": row.dataset_version_id, "title": row.title,
         "status": row.status, "latest_version": row.latest_version_number,
         "updated_at": row.updated_at.isoformat()} for row in rows]}


@router.get("/reports/{report_id}")
def get_report(report_id: int, version: int | None = Query(default=None, ge=1),
               user: User = Depends(current_user), db: Session = Depends(get_db)):
    report = _report(db, report_id, user.id)
    return {"code": 200, "message": "success", "data": _view(report, _version(db, report.id, version))}


@router.get("/artifacts")
def list_artifacts(offset: int = Query(default=0, ge=0), limit: int = Query(default=50, ge=1, le=100),
                   session_id: int | None = Query(default=None, gt=0), type: str | None = None, status: str | None = None,
                   user: User = Depends(current_user), db: Session = Depends(get_db)):
    query = select(AnalysisArtifact).where(AnalysisArtifact.user_id == user.id)
    if session_id:
        query = query.where(AnalysisArtifact.session_id == session_id)
    if type:
        mime_by_type = {'pdf':'application/pdf','python':'text/x-python','sql':'application/sql',
            'excel':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            'word':'application/vnd.openxmlformats-officedocument.wordprocessingml.document'}
        query = query.where(AnalysisArtifact.mime_type == mime_by_type[type] if type in mime_by_type else AnalysisArtifact.kind == type.lower())
    manager = ArtifactManager(db)
    if status:
        # Availability is part of the public status, including missing legacy
        # files with a durable database row. Page the actual matching views.
        views, matched = [], 0
        for item in db.scalars(query.order_by(AnalysisArtifact.created_at.desc(), AnalysisArtifact.id.desc()).execution_options(yield_per=100)):
            view = manager.view(item)
            if view['status'] != status: continue
            if matched >= offset: views.append(view)
            matched += 1
            if len(views) > limit: break
        return {'code':200,'message':'success','data':{'items':views[:limit], 'offset':offset,
            'limit':limit,'has_more':len(views)>limit}}
    rows = db.scalars(query
                      .order_by(AnalysisArtifact.created_at.desc(), AnalysisArtifact.id.desc())
                      .offset(offset).limit(limit + 1)).all()
    page = rows[:limit]
    return {"code": 200, "message": "success", "data": {
        "items": [ArtifactManager(db).view(item) for item in page], "offset": offset,
        "limit": limit, "has_more": len(rows) > limit}}


@router.get("/artifacts/{artifact_id}")
def get_artifact(artifact_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    return {"code": 200, "message": "success", "data": ArtifactManager(db).view(_owned_artifact(db, artifact_id, user.id))}


@router.post('/artifacts/{artifact_id}/regenerations', status_code=201)
def regenerate_artifact(artifact_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    source = _owned_artifact(db,artifact_id,user.id)
    manager = ArtifactManager(db)
    try: row = manager.regenerate(source)
    except LookupError as error:
        raise HTTPException(410,detail={'code':str(error),'message':'原成果不可用，不能重新生成'}) from None
    except ValueError as error:
        raise HTTPException(409,detail={'code':str(error),'message':'成果容量不足，生成失败'}) from None
    return {'code':201,'message':'created','data':manager.view(row)}


@router.get("/artifacts/{artifact_id}/preview")
def preview_artifact(artifact_id: int, offset: int = Query(default=0, ge=0),
                     limit: int = Query(default=100, ge=1, le=200),
                     user: User = Depends(current_user), db: Session = Depends(get_db)):
    artifact = _owned_artifact(db, artifact_id, user.id)
    if expired(artifact):
        raise HTTPException(410, detail={"code": "ARTIFACT_EXPIRED", "message": "文件已过期"})
    if artifact.storage_key:
        try:
            content = LocalArtifactStorage(get_settings().artifact_dir).read(artifact.storage_key)
        except (ValueError, OSError):
            raise HTTPException(410, detail={"code": "ARTIFACT_EXPIRED", "message": "文件已过期"}) from None
        if artifact.mime_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
            from openpyxl import load_workbook
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            sheets = []
            for sheet in book.worksheets:
                rows = sheet.iter_rows(min_row=offset + 1, max_row=offset + limit, values_only=True)
                sheets.append({"name": sheet.title, "rows": [[str(value) if value is not None else None for value in row] for row in rows]})
            book.close()
            data = {"kind": "workbook", "sheets": sheets, "offset": offset, "limit": limit}
        elif artifact.mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            version = db.get(AnalysisReportVersion, artifact.report_version_id)
            data = {"kind": "report_document", "document": version.document_json if version else {}}
        elif artifact.mime_type in {"application/pdf", "text/html", "image/png", "image/svg+xml"}:
            data = {"kind": "document", "mime_type":artifact.mime_type, "preview_url": f"/api/v1/artifacts/{artifact.id}/download?inline=true"}
        else:
            data = {"kind": "text", "content": content[:1024 * 1024].decode("utf-8-sig", errors="replace")}
    else:
        from app.services.artifacts import ArtifactStore
        try:
            payload = ArtifactStore(db).read(artifact)
        except LookupError:
            raise HTTPException(410, detail={"code": "ARTIFACT_EXPIRED", "message": "文件已过期"}) from None
        rows = payload.get("rows", [])[offset:offset + limit]
        data = {"kind": artifact.kind, "columns": payload.get("schema", {}).get("columns", []),
                "rows": rows, "offset": offset, "limit": limit, "total": artifact.row_count}
        if artifact.kind == 'chart': data['chart'] = {**(payload.get('data') or {}),'artifact_id':artifact.id}
    return {"code": 200, "message": "success", "data": {"artifact": ArtifactManager(db).view(artifact), "preview": data}}


@router.patch("/reports/{report_id}")
def update_report(report_id: int, request: ReportRequest, user: User = Depends(current_user),
                  db: Session = Depends(get_db)):
    report = _report(db, report_id, user.id)
    if request.spec.report_id not in (None, report.id):
        raise HTTPException(422, detail={"code": "REPORT_ID_MISMATCH", "message": "报告标识不匹配"})
    try:
        report, version = create_report(db, user.id, request.session_id, request.spec,
                                        request.source_record_ids, report_id=report.id)
    except ReportError as error:
        _error(error)
    return {"code": 200, "message": "success", "data": _view(report, version)}


@router.post("/reports/{report_id}/versions/{version}/exports/{format}", status_code=202)
def export(report_id: int, version: int, format: str, response: Response, request: ExportRequest = ExportRequest(),
           user: User = Depends(current_user), db: Session = Depends(get_db)):
    report = _report(db, report_id, user.id)
    selected = _version(db, report.id, version)
    try:
        record = enqueue_report_task(db, user.id, report.session_id, report.dataset_id,
            report.dataset_version_id, "export", {"report_id": report.id,
            "version_id": selected.id, "format": format}, request_id=request.request_id)
    except ReportError as error:
        _error(error)
    response.status_code = 200 if getattr(record, '_request_replayed', False) else 202
    return {"code": response.status_code, "message": "accepted", "data": {"record_id": record.id,
        "session_id": record.session_id, "status": record.status,
        "status_url": f"/api/v1/analysis/runs/{record.id}",
        "events_url": f"/api/v1/analysis/runs/{record.id}/events"}}

@router.get("/artifacts/{artifact_id}/download")
def download_artifact(artifact_id: int, request: Request, inline: bool = Query(default=False),
                      user: User = Depends(current_user), db: Session = Depends(get_db)):
    artifact = _owned_artifact(db, artifact_id, user.id)
    if expired(artifact):
        raise HTTPException(410, detail={"code": "ARTIFACT_EXPIRED", "message": "文件已过期"})
    try:
        if artifact.storage_key:
            content = LocalArtifactStorage(get_settings().artifact_dir).read(artifact.storage_key)
        else:
            from app.services.artifacts import ArtifactStore
            content = ArtifactStore(db)._path(artifact.stored_name).read_bytes()
    except (ValueError, OSError):
        raise HTTPException(410, detail={"code": "ARTIFACT_EXPIRED", "message": "文件已过期"}) from None
    can_inline = inline and artifact.mime_type in {"application/pdf", "text/html", "image/png", "image/svg+xml"}
    disposition = "inline" if can_inline else "attachment"
    encoded_name = quote(artifact.file_name or f"artifact-{artifact.id}.json")
    headers = {"Content-Disposition": f"{disposition}; filename*=UTF-8''{encoded_name}",
               "X-Content-Type-Options": "nosniff", "Cache-Control": "private, no-store",
               "Accept-Ranges": "bytes"}
    status_code = 200
    byte_range = request.headers.get("range")
    if byte_range:
        import re
        match = re.fullmatch(r"bytes=(\d*)-(\d*)", byte_range.strip())
        if not match or (not match.group(1) and not match.group(2)):
            raise HTTPException(416, detail={"code": "RANGE_NOT_SATISFIABLE", "message": "文件范围无效"},
                                headers={"Content-Range": f"bytes */{len(content)}"})
        if match.group(1):
            start = int(match.group(1)); end = int(match.group(2) or len(content) - 1)
        else:
            length = int(match.group(2)); start, end = max(0, len(content) - length), len(content) - 1
        if start >= len(content) or end < start:
            raise HTTPException(416, detail={"code": "RANGE_NOT_SATISFIABLE", "message": "文件范围无效"},
                                headers={"Content-Range": f"bytes */{len(content)}"})
        end = min(end, len(content) - 1)
        content = content[start:end + 1]
        headers["Content-Range"] = f"bytes {start}-{end}/{artifact.size_bytes}"
        status_code = 206
    return Response(content=content, status_code=status_code,
                    media_type=artifact.mime_type or "application/json", headers=headers)
