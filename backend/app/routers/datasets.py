from __future__ import annotations

import logging
import os
import tempfile
import uuid
from datetime import UTC, datetime
from pathlib import Path

from fastapi import APIRouter, BackgroundTasks, Depends, File, HTTPException, Query, UploadFile, Request, Form
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.config import get_settings
from app.database import SessionLocal, engine, get_db
from app.dependencies import current_user
from app.models import Dataset, DatasetColumn, User, BackgroundJob, AnalysisRecord, AnalysisArtifact, CleanupTask
from app.services.datasets import drop_projection, process_dataset, projection_table


logger = logging.getLogger(__name__)
router = APIRouter(prefix="/datasets", tags=["datasets"])
settings = get_settings()
CHUNK_BYTES = 1024 * 1024


def _error(status_code: int, code: str, message: str) -> HTTPException:
    return HTTPException(status_code=status_code, detail={"code": code, "message": message})


def _owned_dataset(db: Session, dataset_id: int, user_id: int) -> Dataset:
    dataset = db.get(Dataset, dataset_id)
    if dataset is None:
        raise _error(404, "DATASET_NOT_FOUND", "数据集不存在")
    if dataset.user_id != user_id:
        raise _error(403, "DATASET_FORBIDDEN", "你没有权限访问这个数据集")
    return dataset


def _iso(value: datetime | None) -> str | None:
    if value is None:
        return None
    return value.replace(tzinfo=UTC).isoformat() if value.tzinfo is None else value.isoformat()


def _summary(dataset: Dataset) -> dict:
    return {
        "id": dataset.id,
        "uploaded_file_id": dataset.uploaded_file_id,
        "origin_metadata": dataset.origin_metadata_json,
        "original_name": dataset.original_name,
        "file_type": dataset.file_type,
        "file_size": dataset.file_size,
        "row_count": dataset.row_count,
        "column_count": dataset.column_count,
        "status": dataset.status,
        "parse_error_code": dataset.parse_error_code,
        "parse_error_message": dataset.parse_error_message,
        "quality_warnings": dataset.quality_warnings_json or [],
        "parse_version": dataset.parse_version or "1",
        "current_version_id": dataset.current_version_id,
        "created_at": _iso(dataset.created_at),
        "updated_at": _iso(dataset.updated_at),
    }


def _safe_original_name(name: str | None) -> str:
    cleaned = (name or "dataset").replace("\\", "/").split("/")[-1].strip()
    return cleaned[:255] or "dataset"


def _json_value(value):
    if hasattr(value, "isoformat"):
        return value.isoformat()
    if hasattr(value, "item"):
        value = value.item()
    return value


@router.post("/inspect-sheets")
def inspect_workbook_sheets(file: UploadFile = File(...), user: User = Depends(current_user)):
    suffix = Path(_safe_original_name(file.filename)).suffix.lower()
    if suffix not in {".xls", ".xlsx"}:
        raise _error(415, "WORKSHEET_SELECTION_UNAVAILABLE", "仅 Excel 文件支持工作表选择")
    content = file.file.read(settings.max_upload_bytes + 1)
    if len(content) > settings.max_upload_bytes:
        raise _error(413, "DATASET_FILE_TOO_LARGE", "文件超过 20 MB 限制")
    if not content:
        raise _error(400, "DATASET_EMPTY", "请选择非空文件")
    path = None
    try:
        with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as temporary:
            temporary.write(content)
            path = Path(temporary.name)
        if suffix == ".xlsx":
            from openpyxl import load_workbook
            from app.services.datasets import MAX_COMPRESSION_RATIO, MAX_XLSX_UNCOMPRESSED_BYTES
            import zipfile
            with zipfile.ZipFile(path) as archive:
                infos = archive.infolist()
                expanded = sum(item.file_size for item in infos)
                compressed = max(1, sum(item.compress_size for item in infos))
                if expanded > MAX_XLSX_UNCOMPRESSED_BYTES or expanded / compressed > MAX_COMPRESSION_RATIO:
                    raise _error(413, "DATASET_ARCHIVE_TOO_LARGE", "Excel 文件解压后过大，无法安全解析")
                if any(item.filename.lower().endswith("vbaproject.bin") for item in infos):
                    raise _error(415, "DATASET_MACRO_UNSUPPORTED", "不支持包含宏的 Excel 文件")
            workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
            sheets = [sheet.title for sheet in workbook.worksheets
                      if sum(any(value is not None for value in row)
                             for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row or 1, 3), values_only=True)) >= 2]
            workbook.close()
        else:
            import pandas as pd
            sheets = pd.ExcelFile(path, engine="xlrd").sheet_names
        if not sheets:
            raise _error(400, "DATASET_EMPTY", "Excel 文件中没有非空工作表")
        return {"code": 200, "message": "success", "data": {"sheets": sheets, "default": sheets[0]}}
    except HTTPException:
        raise
    except Exception:
        raise _error(400, "DATASET_PARSE_FAILED", "无法读取 Excel 工作表") from None
    finally:
        if path:
            path.unlink(missing_ok=True)


@router.post("/upload", status_code=202)
def upload_dataset(
    request: Request,
    background_tasks: BackgroundTasks,
    file: UploadFile = File(...),
    sheet_name: str | None = Form(default=None, max_length=128),
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
    session_id: int | None = Form(default=None),
):
    original_name = _safe_original_name(file.filename)
    extension = Path(original_name).suffix.lower().lstrip(".")
    if extension not in {"csv", "tsv", "json", "jsonl", "xlsx", "xls", "parquet"}:
        raise _error(415, "DATASET_TYPE_NOT_SUPPORTED", "支持 CSV、TSV、JSON、XLS、XLSX 和 Parquet 文件")

    upload_dir = Path(settings.upload_dir)
    upload_dir.mkdir(parents=True, exist_ok=True)
    stored_name = f"{uuid.uuid4().hex}.{extension}"
    temporary_path = upload_dir / f"{stored_name}.part"
    final_path = upload_dir / stored_name
    file_size = 0
    try:
        with temporary_path.open("wb") as output:
            while chunk := file.file.read(CHUNK_BYTES):
                file_size += len(chunk)
                if file_size > settings.max_upload_bytes:
                    raise _error(413, "DATASET_FILE_TOO_LARGE", "文件超过 20 MB 限制")
                output.write(chunk)
        if file_size == 0:
            raise _error(400, "DATASET_EMPTY", "请选择非空文件")
        if extension in {'csv', 'tsv', 'json', 'jsonl'}:
            from app.files.parsers import validate_signature, FileParseError
            try:
                validate_signature(temporary_path, extension)
            except FileParseError as exc:
                raise _error(400, exc.code, exc.message) from None
        if extension in {"xlsx", "parquet"}:
            with temporary_path.open("rb") as source:
                signature = source.read(4)
                if extension == "xlsx" and signature != b"PK\x03\x04":
                    raise _error(400, "DATASET_PARSE_FAILED", "XLSX 文件签名不正确")
                if extension == "parquet" and signature != b"PAR1":
                    raise _error(400, "DATASET_PARSE_FAILED", "Parquet 文件签名不正确")
        if extension == "xls":
            with temporary_path.open("rb") as source:
                if source.read(8) != bytes.fromhex("D0CF11E0A1B11AE1"):
                    raise _error(400, "DATASET_PARSE_FAILED", "XLS 文件签名不正确")
        os.replace(temporary_path, final_path)
        now = datetime.now(UTC)
        dataset = Dataset(
            user_id=user.id,
            original_name=original_name,
            stored_name=stored_name,
            file_type=extension,
            file_size=file_size,
            status="parsing",
            parse_options_json={"sheet_name": sheet_name} if sheet_name else None,
            created_at=now,
            updated_at=now,
        )
        from app.services.analysis import resource_lock
        with resource_lock:
            from app.files.service import register_dataset_file
            register_dataset_file(db, dataset, final_path, file.content_type)
            db.add(dataset)
            db.flush()
            from app.files.service import bind_session
            bind_session(db, session_id, user.id, dataset.uploaded_file_id, dataset.id)
            if getattr(request.app.state, "task_supervisor", None):
                from app.database import projection_engine
                dataset.projection_schema = projection_engine.url.database
                dataset.projection_table = f"dataset_{dataset.id}"
                queued = db.scalar(select(func.count()).select_from(BackgroundJob).where(BackgroundJob.status.in_(("pending", "running")))) or 0
                if queued >= settings.max_pending_jobs:
                    raise _error(429, "TASK_QUEUE_FULL", "任务队列已满，请稍后重试")
                db.add(BackgroundJob(kind="parse", resource_id=dataset.id, dataset_id=dataset.id, user_id=user.id, status="pending", created_at=now))
            db.commit()
            db.refresh(dataset)
    except HTTPException:
        db.rollback()
        temporary_path.unlink(missing_ok=True)
        final_path.unlink(missing_ok=True)
        raise
    except Exception:
        db.rollback()
        temporary_path.unlink(missing_ok=True)
        final_path.unlink(missing_ok=True)
        logger.exception("Dataset upload failed user_id=%s", user.id)
        raise _error(500, "DATASET_UPLOAD_FAILED", "文件上传失败，请重试")
    finally:
        file.file.close()

    if not getattr(request.app.state, "task_supervisor", None):
        # Embedded app without lifespan keeps the original background-task contract.
        background_tasks.add_task(process_dataset, dataset.id, SessionLocal, engine, settings.upload_dir)
    return {"code": 202, "message": "parsing", "data": _summary(dataset)}


@router.get("")
def list_datasets(
    q: str | None = None,
    file_type: str | None = Query(default=None, pattern="^(csv|xlsx)$"),
    status_filter: str | None = Query(default=None, alias="status"),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=10, ge=1, le=50),
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    query = select(Dataset).where(Dataset.user_id == user.id)
    count_query = select(func.count()).select_from(Dataset).where(Dataset.user_id == user.id)
    if q and q.strip():
        pattern = f"%{q.strip()}%"
        query = query.where(Dataset.original_name.ilike(pattern))
        count_query = count_query.where(Dataset.original_name.ilike(pattern))
    if file_type:
        query = query.where(Dataset.file_type == file_type)
        count_query = count_query.where(Dataset.file_type == file_type)
    if status_filter:
        query = query.where(Dataset.status == status_filter)
        count_query = count_query.where(Dataset.status == status_filter)
    total = db.scalar(count_query) or 0
    items = db.scalars(query.order_by(Dataset.created_at.desc()).offset((page - 1) * page_size).limit(page_size)).all()
    return {
        "code": 200,
        "message": "success",
        "data": {"items": [_summary(item) for item in items], "page": page, "page_size": page_size, "total": total,
                 "pages": (total + page_size - 1) // page_size},
    }


@router.get("/{dataset_id}")
def get_dataset(dataset_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    dataset = _owned_dataset(db, dataset_id, user.id)
    return {"code": 200, "message": "success", "data": _summary(dataset)}


@router.get("/{dataset_id}/columns")
def get_columns(dataset_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    _owned_dataset(db, dataset_id, user.id)
    columns = db.scalars(
        select(DatasetColumn).where(DatasetColumn.dataset_id == dataset_id).order_by(DatasetColumn.ordinal_position)
    ).all()
    return {
        "code": 200,
        "message": "success",
        "data": [
            {"name": column.name, "original_name": column.original_name, "data_type": column.data_type,
             "nullable": column.nullable, "missing_count": column.missing_count, "unique_count": column.unique_count,
             "sample_values": column.sample_values_json or []}
            for column in columns
        ],
    }


@router.get("/{dataset_id}/preview")
def preview_dataset(
    dataset_id: int,
    offset: int = Query(default=0, ge=0),
    limit: int = Query(default=50, ge=1, le=100),
    columns: str | None = None,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    dataset = _owned_dataset(db, dataset_id, user.id)
    if dataset.status != "ready":
        raise _error(409, "DATASET_NOT_READY", "数据集尚未完成解析")
    try:
        from app.database import projection_engine
        from app.services.analysis import select_projection_bind
        bind = select_projection_bind(dataset, engine, projection_engine)
        table = projection_table(dataset_id, bind, dataset.projection_table)
    except LookupError:
        raise _error(409, "DATASET_NOT_READY", "数据预览暂不可用，请重新上传数据集")
    column_names = [column.name for column in db.scalars(
        select(DatasetColumn).where(DatasetColumn.dataset_id == dataset_id).order_by(DatasetColumn.ordinal_position)
    ).all()]
    requested = [item.strip() for item in columns.split(",") if item.strip()] if columns else column_names
    if any(name not in column_names for name in requested):
        raise _error(422, "DATASET_UNKNOWN_COLUMN", "预览字段不存在")
    selected = [table.c[name] for name in requested]
    with bind.connect() as connection:
        rows = connection.execute(select(*selected).select_from(table).offset(offset).limit(limit)).mappings().all()
    return {
        "code": 200,
        "message": "success",
        "data": {"columns": requested, "rows": [{key: _json_value(value) for key, value in row.items()} for row in rows],
                 "offset": offset, "limit": limit, "total_rows": dataset.row_count or 0},
    }


@router.delete("/{dataset_id}")
def delete_dataset(dataset_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    dataset = _owned_dataset(db, dataset_id, user.id)
    if dataset.status in {"uploading", "parsing"}:
        raise _error(409, "DATASET_BUSY", "数据集仍在解析，请稍后删除")
    db.scalar(select(Dataset).where(Dataset.id == dataset_id).with_for_update())
    if db.scalar(select(BackgroundJob.id).where(BackgroundJob.dataset_id==dataset_id,BackgroundJob.kind=='tool',BackgroundJob.status.in_(('pending','running'))).limit(1)):
        raise _error(409,'DATASET_BUSY','数据集正在执行工具任务')
    from app.models import ToolExecutionRecord
    if any(record.parameters_json.get('right_dataset_id')==dataset_id for record in db.scalars(select(ToolExecutionRecord).where(ToolExecutionRecord.tool_name=='publish_join',ToolExecutionRecord.status.in_(('pending','running'))))):
        raise _error(409,'DATASET_BUSY','数据集正在作为关联任务的输入')
    if db.scalar(select(AnalysisRecord.id).where(AnalysisRecord.dataset_id == dataset_id, AnalysisRecord.status.in_(("pending", "running"))).limit(1)):
        raise _error(409, "DATASET_BUSY", "数据集正在分析，请稍后删除")
    stored_name = dataset.stored_name
    from app.models import DatasetVersion
    projections=[{'table':version.projection_table,'schema':version.projection_schema} for version in db.scalars(select(DatasetVersion).where(DatasetVersion.dataset_id==dataset_id))]
    cleanup = CleanupTask(payload_json={"dataset_id": dataset_id, "stored_name": stored_name,
        "projection_schema": dataset.projection_schema,
        "artifacts": list(db.scalars(select(AnalysisArtifact.stored_name).where(AnalysisArtifact.dataset_id == dataset_id))),
        "artifact_storage_keys": list(db.scalars(select(AnalysisArtifact.storage_key).where(
            AnalysisArtifact.dataset_id == dataset_id, AnalysisArtifact.storage_key.is_not(None))))},
        status="pending", created_at=datetime.now(UTC))
    db.add(cleanup)
    cleanup.payload_json=dict(cleanup.payload_json,projections=projections)
    # Legacy table uploads are one source per dataset. Document extractions own
    # a separate generated CSV, so their shared original remains registered.
    from app.models import UploadedFile
    source = db.get(UploadedFile, dataset.uploaded_file_id) if dataset.uploaded_file_id else None
    if source and source.stored_name == stored_name and not db.scalar(select(Dataset.id).where(Dataset.uploaded_file_id == source.id, Dataset.id != dataset_id).limit(1)):
        from app.models import AnalysisSession
        for session in db.scalars(select(AnalysisSession).where(AnalysisSession.user_id == user.id).with_for_update()):
            context = dict(session.context_json or {})
            if source.id in context.get('uploaded_file_ids', []):
                context['uploaded_file_ids'] = [fid for fid in context['uploaded_file_ids'] if fid != source.id]
                session.context_json = context
        dataset.uploaded_file_id = None
        db.flush()
        db.delete(source)
    db.delete(dataset)
    db.commit()
    from app.services.jobs import perform_cleanup
    from app.database import projection_engine
    perform_cleanup(db, cleanup, engine, projection_engine)
    return {"code": 200, "message": "success", "data": {"deleted_id": dataset_id, "deleted_at": datetime.now(UTC).isoformat()}}
