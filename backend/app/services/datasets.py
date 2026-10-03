from __future__ import annotations

import csv
import hashlib
import logging
import math
import re
import zipfile
import sys
from datetime import UTC, date, datetime, time
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from fastapi import HTTPException
from sqlalchemy import BigInteger, Boolean, Date, DateTime, MetaData, Numeric, Table, Text, inspect, select
from sqlalchemy.orm import Session, sessionmaker

from app.database import engine, SessionLocal
from app.models import Dataset, DatasetColumn, DatasetVersion
from app.config import get_settings


logger = logging.getLogger(__name__)
MAX_XLSX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_COMPRESSION_RATIO = 100


class DatasetParseError(Exception):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.public_message = message


def safe_column_names(headers: list[Any]) -> list[str]:
    names: list[str] = []
    raw_names: set[str] = set()
    for index, header in enumerate(headers, start=1):
        text = "" if header is None else str(header).strip()
        if text and text in raw_names:
            raise DatasetParseError("DATASET_DUPLICATE_COLUMNS", "文件包含重复字段名，请修改表头后重新上传")
        raw_names.add(text)
        name = re.sub(r"[^A-Za-z0-9_]", "_", text).strip("_").lower()
        if not name or not re.match(r"^[a-z_]", name):
            name = f"column_{index}"
        if len(name) > 64 or name in names:
            suffix = hashlib.sha256((text + ':' + str(index)).encode()).hexdigest()[:10]
            name = name[:53] + '_' + suffix
        names.append(name)
    return names


def _read_csv(path: Path) -> pd.DataFrame:
    last_error: Exception | None = None
    for encoding in ("utf-8-sig", "utf-8"):
        try:
            with path.open("r", encoding=encoding, newline="") as handle:
                header = next(csv.reader(handle), None)
            if not header:
                raise DatasetParseError("DATASET_EMPTY", "文件为空或缺少表头")
            safe_column_names(header)
            if len(header) > get_settings().max_dataset_columns:
                raise DatasetParseError("DATASET_TOO_MANY_COLUMNS", f"文件超过 {get_settings().max_dataset_columns} 列限制")
            chunks = []
            count = 0
            with pd.read_csv(path, encoding=encoding, dtype='string', chunksize=5000) as reader:
                for chunk in reader:
                    count += len(chunk)
                    if count > get_settings().max_dataset_rows:
                        raise DatasetParseError("DATASET_TOO_MANY_ROWS", f"文件超过 {get_settings().max_dataset_rows} 行限制")
                    chunks.append(chunk)
                    if sum(int(part.memory_usage(deep=True).sum()) for part in chunks) > get_settings().dataframe_max_bytes:
                        raise DatasetParseError("DATASET_MEMORY_LIMIT", "文件超过解析内存预算")
            frame = pd.concat(chunks, ignore_index=True) if chunks else pd.DataFrame(columns=header)
            frame.attrs['transformations'] = [{'operation': 'csv_read', 'encoding': encoding}, {'operation': 'numeric_inference', 'columns': []}]
            for name in frame.columns:
                values = frame[name].dropna()
                if not values.empty and not values.str.match(r'^0\d+').any():
                    numeric = pd.to_numeric(frame[name], errors='coerce')
                    if numeric.notna().sum() == values.size:
                        frame[name] = numeric
                        frame.attrs['transformations'][1]['columns'].append(str(name))
            if frame.empty or len(frame.columns) == 0:
                raise DatasetParseError("DATASET_EMPTY", "文件没有可解析的数据行")
            return frame
        except UnicodeDecodeError as exc:
            last_error = exc
        except pd.errors.EmptyDataError as exc:
            raise DatasetParseError("DATASET_EMPTY", "文件为空或缺少表头") from exc
        except pd.errors.ParserError as exc:
            raise DatasetParseError("DATASET_PARSE_FAILED", "CSV 文件格式不正确") from exc
    raise DatasetParseError("DATASET_ENCODING_UNSUPPORTED", "CSV 编码无法识别，请另存为 UTF-8 后重试") from last_error


def _read_xlsx(path: Path) -> pd.DataFrame:
    try:
        with zipfile.ZipFile(path) as archive:
            infos = archive.infolist()
            expanded = sum(entry.file_size for entry in infos)
            compressed = max(1, sum(entry.compress_size for entry in infos))
            if expanded > MAX_XLSX_UNCOMPRESSED_BYTES or expanded / compressed > MAX_COMPRESSION_RATIO:
                raise DatasetParseError("DATASET_ARCHIVE_TOO_LARGE", "Excel 文件解压后过大，无法安全解析")
            if any(entry.filename.lower().endswith("vbaproject.bin") for entry in infos):
                raise DatasetParseError("DATASET_MACRO_UNSUPPORTED", "不支持包含宏的 Excel 文件")
    except zipfile.BadZipFile as exc:
        raise DatasetParseError("DATASET_PARSE_FAILED", "Excel 文件内容损坏") from exc

    workbook = None
    try:
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        for sheet in workbook.worksheets:
            if sheet.max_column and sheet.max_column > get_settings().max_dataset_columns:
                raise DatasetParseError("DATASET_TOO_MANY_COLUMNS", f"文件超过 {get_settings().max_dataset_columns} 列限制")
            rows: list[tuple[Any, ...]] = []
            memory = 0
            for row in sheet.iter_rows(values_only=True):
                if any(value is not None for value in row):
                    rows.append(row)
                    memory += sys.getsizeof(row) + sum(sys.getsizeof(value) for value in row)
                    if memory > get_settings().dataframe_max_bytes:
                        raise DatasetParseError("DATASET_MEMORY_LIMIT", "文件超过解析内存预算")
                if len(rows) > get_settings().max_dataset_rows + 1:
                    raise DatasetParseError("DATASET_TOO_MANY_ROWS", f"文件超过 {get_settings().max_dataset_rows} 行限制")
            if rows:
                headers = list(rows[0])
                width = len(headers)
                frame = pd.DataFrame([list(row[:width]) for row in rows[1:]], columns=headers,dtype=object).convert_dtypes()
                frame.attrs['transformations'] = [{'operation': 'worksheet_selection', 'sheet': sheet.title}, {'operation': 'drop_empty_rows', 'count': max(0, (sheet.max_row or len(rows)) - len(rows))}]
                frame = frame.dropna(how="all")
                formula_book = None
                try:
                    formula_book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
                    missing_cache = 0
                    formatted_identifiers = 0
                    for raw, cached in zip(formula_book[sheet.title].iter_rows(), sheet.iter_rows(values_only=True)):
                        missing_cache += sum(cell.data_type == "f" and value is None for cell, value in zip(raw, cached))
                        formatted_identifiers += sum(cell.data_type == "n" and bool(re.fullmatch(r"0{2,}", cell.number_format or "")) for cell in raw)
                    if missing_cache:
                        frame.attrs['quality_warnings'] = [{'code': 'FORMULA_CACHE_MISSING', 'count': missing_cache, 'message': '部分 Excel 公式没有缓存结果；请在 Excel 重新计算并保存'}]
                    if formatted_identifiers:
                        frame.attrs.setdefault('quality_warnings', []).append({'code': 'EXCEL_DISPLAY_FORMAT', 'count': formatted_identifiers, 'message': 'Excel 数字显示格式中的前导零不会作为原始字符保留；编号建议存为文本'})
                finally:
                    if formula_book is not None:
                        formula_book.close()
                if frame.empty:
                    raise DatasetParseError("DATASET_EMPTY", "工作表只有表头，没有数据行")
                return frame
        raise DatasetParseError("DATASET_EMPTY", "Excel 文件中没有非空工作表")
    except DatasetParseError:
        raise
    except Exception as exc:
        raise DatasetParseError("DATASET_PARSE_FAILED", "Excel 工作表无法读取") from exc
    finally:
        if workbook is not None:
            workbook.close()


def parse_file(path: Path, file_type: str) -> pd.DataFrame:
    frame = _read_csv(path) if file_type == "csv" else _read_xlsx(path)
    if len(frame) > get_settings().max_dataset_rows:
        raise DatasetParseError("DATASET_TOO_MANY_ROWS", f"文件超过 {get_settings().max_dataset_rows} 行限制")
    if len(frame.columns) > get_settings().max_dataset_columns:
        raise DatasetParseError("DATASET_TOO_MANY_COLUMNS", f"文件超过 {get_settings().max_dataset_columns} 列限制")
    if int(frame.memory_usage(deep=True).sum()) > get_settings().dataframe_max_bytes:
        raise DatasetParseError('DATASET_MEMORY_LIMIT', '文件超过解析内存预算')
    original_names = list(frame.columns)
    frame.attrs["original_columns"] = [str(name) for name in original_names]
    frame.columns = safe_column_names(original_names)
    frame.attrs.setdefault('transformations', []).append({'operation': 'normalize_columns', 'mapping': [{'original': str(original), 'normalized': str(normalized)} for original, normalized in zip(original_names, frame.columns)]})
    before_types = {str(name): str(frame[name].dtype) for name in frame.columns}
    frame = _infer_date_columns(frame)
    frame.attrs['transformations'].append({'operation': 'date_inference', 'changes': [{'column': name, 'from': before_types[name], 'to': str(frame[name].dtype)} for name in frame.columns if before_types[name] != str(frame[name].dtype)]})
    if any(pd.api.types.is_float_dtype(frame[name].dtype) for name in frame.columns):
        frame.attrs.setdefault('quality_warnings', []).append({'code': 'NUMERIC_PRECISION', 'message': '数值计算使用浮点近似，投影小数最多保留8位；本应用不用于财务结算'})
    empty_count = int(frame.isna().all(axis=1).sum())
    frame.attrs['transformations'].append({'operation': 'drop_empty_rows', 'count': empty_count})
    frame = frame.dropna(how="all").reset_index(drop=True)
    if frame.empty:
        raise DatasetParseError("DATASET_EMPTY", "文件没有可解析的数据行")
    return frame


def _infer_date_columns(frame: pd.DataFrame) -> pd.DataFrame:
    frame.attrs.setdefault('quality_warnings', [])
    for name in frame.columns:
        series = frame[name]
        if not (pd.api.types.is_object_dtype(series.dtype) or pd.api.types.is_string_dtype(series.dtype)):
            continue
        values = series.dropna()
        if values.empty:
            continue
        sample = values.head(20).astype(str)
        name_hint = any(part in name.lower() for part in ("date", "time", "timestamp"))
        value_hint = sample.str.match(r"^\d{4}[-/]\d{1,2}[-/]\d{1,2}").any()
        if not name_hint and not value_hint:
            continue
        parsed = pd.to_datetime(series, errors="coerce")
        failed = int((series.notna() & parsed.isna()).sum())
        if failed:
            frame.attrs['quality_warnings'].append({'code': 'DATE_PARSE_FAILED', 'column': name, 'count': failed, 'message': '保留原始字段值；部分日期无法解析'})
            continue
        if not parsed.notna().any():
            continue
        valid_times = parsed.dropna().dt.time
        if not valid_times.empty and valid_times.eq(time.min).all():
            frame[name] = parsed.dt.date
        else:
            frame[name] = parsed
    return frame


def _data_type(series: pd.Series) -> tuple[str, Any]:
    if pd.api.types.is_bool_dtype(series.dtype):
        return "boolean", Boolean()
    if pd.api.types.is_integer_dtype(series.dtype):
        return "integer", BigInteger()
    if pd.api.types.is_numeric_dtype(series.dtype):
        return "decimal", Numeric(24, 8)
    if pd.api.types.is_datetime64_any_dtype(series.dtype):
        return "datetime", DateTime()
    if series.dropna().map(lambda value: isinstance(value, date) and not isinstance(value, datetime)).all() and series.notna().any():
        return "date", Date()
    return "string", Text()


def _sample_value(value: Any) -> Any:
    if hasattr(value, "item"):
        value = value.item()
    if pd.isna(value):
        return None
    if isinstance(value, (pd.Timestamp, datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _sql_value(value: Any) -> Any:
    if pd.isna(value):
        return None
    if hasattr(value, "item"):
        return value.item()
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    return value


def write_projection(frame: pd.DataFrame, dataset_id: int, bind=engine, table_name=None) -> str:
    table_name = table_name or f"dataset_{int(dataset_id)}"
    if not re.fullmatch(rf'dataset_{int(dataset_id)}(?:_v_[a-f0-9]{{32}})?',table_name):
        raise ValueError('invalid server projection name')
    dtype: dict[str, Any] = {}
    for name in frame.columns:
        _, sql_type = _data_type(frame[name])
        dtype[name] = sql_type
    # Mapping nullable BIGINT through inferred float columns rounds values above
    # 2**53. Keep Python int/None scalars until the SQL driver binds them.
    clean = pd.DataFrame({name:pd.Series([_sql_value(value) for value in frame[name]],dtype=object) for name in frame},index=frame.index)
    clean.to_sql(table_name, con=bind, if_exists="fail", index=False, dtype=dtype)
    return table_name


def profile_frame(frame: pd.DataFrame, dataset_id: int) -> list[DatasetColumn]:
    now = datetime.now(UTC)
    columns: list[DatasetColumn] = []
    original_names = frame.attrs.get("original_columns", list(frame.columns))
    for position, name in enumerate(frame.columns):
        series = frame[name]
        type_name, _ = _data_type(series)
        samples = [_sample_value(value) for value in series.dropna().head(5).tolist()]
        original = str(original_names[position])
        columns.append(
            DatasetColumn(
                dataset_id=dataset_id,
                ordinal_position=position,
                name=str(name),
                original_name=original,
                data_type=type_name,
                nullable=bool(series.isna().any()),
                missing_count=int(series.isna().sum()),
                unique_count=int(series.nunique(dropna=True)),
                sample_values_json=samples,
                created_at=now,
            )
        )
    return columns


def process_dataset(
    dataset_id: int,
    session_factory: sessionmaker = SessionLocal,
    bind=engine,
    upload_dir: str | None = None,
    lease_guard=None,
) -> None:
    with session_factory() as db:
        dataset = db.get(Dataset, dataset_id)
        if dataset is None or dataset.current_version_id is not None:
            return
        if lease_guard and not lease_guard(db, False):
            db.rollback()
            return
        dataset.status = "parsing"
        dataset.updated_at = datetime.now(UTC)
        db.commit()
        # stored_name is server-generated; never accept a path from the request.
        if upload_dir is None:
            from app.config import get_settings

            upload_dir = get_settings().upload_dir
        path = Path(upload_dir) / dataset.stored_name
        try:
            frame = parse_file(path, dataset.file_type)
            if lease_guard and not lease_guard(db, False):
                db.rollback()
                return
            db.commit()
            write_projection(frame, dataset_id, bind=bind)
            db.query(DatasetColumn).filter(DatasetColumn.dataset_id == dataset_id).delete()
            db.add_all(profile_frame(frame, dataset_id))
            dataset.row_count = int(len(frame))
            dataset.column_count = int(len(frame.columns))
            if hasattr(dataset, 'quality_warnings_json'):
                dataset.quality_warnings_json = frame.attrs.get('quality_warnings', [])
                dataset.parse_version = '3'
                dataset.projection_table = f'dataset_{int(dataset_id)}'
                dataset.projection_schema = getattr(getattr(bind, 'url', None), 'database', None)
            dataset.status = "ready"
            dataset.parse_error_code = None
            dataset.parse_error_message = None
            dataset.updated_at = datetime.now(UTC)
            if lease_guard and not lease_guard(db, True):
                db.rollback()
                drop_projection(dataset_id, bind)
                return
            from app.datasets.versions import create_initial_version
            create_initial_version(db, dataset, frame, upload_dir)
            db.commit()
        except DatasetParseError as exc:
            db.rollback()
            if lease_guard and not lease_guard(db, True):
                db.rollback()
                return
            try:
                drop_projection(dataset_id, bind)
            except Exception:
                logger.error("Could not remove failed dataset projection dataset_id=%s", dataset_id)
            _mark_parse_failed(db, dataset_id, exc.code, exc.public_message)

        except Exception:
            logger.error("Dataset parsing failed dataset_id=%s", dataset_id)
            db.rollback()
            if lease_guard and not lease_guard(db, True):
                db.rollback()
                return
            try:
                drop_projection(dataset_id, bind)
            except Exception:
                logger.error("Could not remove failed dataset projection dataset_id=%s", dataset_id)
            _mark_parse_failed(db, dataset_id, "DATASET_PARSE_FAILED", "文件无法解析，请检查格式后重新上传")



def _mark_parse_failed(db: Session, dataset_id: int, code: str, message: str) -> None:
    dataset = db.get(Dataset, dataset_id)
    if dataset is not None:
        dataset.status = "failed"
        dataset.parse_error_code = code
        dataset.parse_error_message = message
        dataset.updated_at = datetime.now(UTC)
        db.commit()


def recover_interrupted_datasets(
    session_factory: sessionmaker = SessionLocal,
    bind=engine,
    upload_dir: str | None = None,
) -> int:
    """Fail interrupted uploads, retain accepted sources and release unpublished projections."""
    if upload_dir is None:
        from app.config import get_settings

        upload_dir = get_settings().upload_dir
    recovered: list[tuple[int, str]] = []
    with session_factory() as db:
        pending = db.query(Dataset).filter(Dataset.status.in_(("uploading", "parsing"))).all()
        now = datetime.now(UTC)
        for dataset in pending:
            dataset.status = "failed"
            dataset.parse_error_code = "DATASET_PROCESS_INTERRUPTED"
            dataset.parse_error_message = "服务重启中断了解析，请重新上传文件"
            dataset.updated_at = now
            recovered.append((dataset.id, dataset.stored_name))
        if recovered:
            db.commit()
    for dataset_id, stored_name in recovered:
        try:
            drop_projection(dataset_id, bind)
        except Exception:
            logger.error("Could not clean interrupted dataset projection dataset_id=%s", dataset_id)

    return len(recovered)


def projection_table(dataset_id: int, bind=engine, table_name=None) -> Table:
    name = table_name or f"dataset_{int(dataset_id)}"
    if not re.fullmatch(rf'dataset_{int(dataset_id)}(?:_v_[a-f0-9]{{32}})?',name):
        raise ValueError('invalid server projection name')
    if not inspect(bind).has_table(name):
        raise LookupError("Dataset projection is unavailable")
    return Table(name, MetaData(), autoload_with=bind)


def drop_projection(dataset_id: int, bind=engine) -> None:
    name = f"dataset_{int(dataset_id)}"
    if inspect(bind).has_table(name):
        Table(name, MetaData(), autoload_with=bind).drop(bind)


class DatasetService:
    """One authorized metadata and typed projection loading boundary."""

    def __init__(self, db: Session, projection_bind=engine):
        self.db = db
        self.projection_bind = projection_bind

    def get(self, dataset_id: int, user_id: int, ready: bool = True) -> tuple[Dataset, list[DatasetColumn]]:
        dataset = self.db.query(Dataset).filter(Dataset.id == dataset_id, Dataset.user_id == user_id).first()
        if dataset is None:
            raise HTTPException(status_code=404, detail={'code': 'DATASET_NOT_FOUND', 'message': '数据集不存在'})
        if ready and dataset.status != 'ready':
            raise HTTPException(status_code=409, detail={'code': 'DATASET_NOT_READY', 'message': '数据集尚未解析完成'})
        columns = self.db.query(DatasetColumn).filter(DatasetColumn.dataset_id == dataset.id).order_by(DatasetColumn.ordinal_position).all()
        return dataset, columns

    @staticmethod
    def metadata(dataset: Dataset, columns: list[DatasetColumn]) -> dict[str, Any]:
        return {'dataset_version_id': getattr(dataset, 'current_version_id', None), 'dataset_id': dataset.id, 'row_count': dataset.row_count, 'columns': [{'name': column.name, 'label': column.original_name, 'data_type': column.data_type, 'missing_count': column.missing_count} for column in columns], 'quality_warnings': getattr(dataset, 'quality_warnings_json', None) or [], 'parse_version': getattr(dataset, 'parse_version', None) or '1'}

    def model_metadata(self, dataset: Dataset, columns: list[DatasetColumn], version_id: int | None = None) -> dict[str, Any]:
        result = self.metadata(dataset, columns)
        version = self.get_version(dataset, version_id)
        if version is not None:
            result['dataset_version_id'] = version.id
            from app.datasets.schemas import DatasetSchema, DatasetProfile
            from app.datasets.profiler import model_summary
            result['profile'] = model_summary(DatasetSchema.model_validate(version.schema_json), DatasetProfile.model_validate(version.profile_json))
        return result

    def get_version(self, dataset: Dataset, version_id: int | None = None) -> DatasetVersion | None:
        identifier = version_id if version_id is not None else dataset.current_version_id
        if identifier is None:
            return None
        version = self.db.get(DatasetVersion, identifier)
        if version is None or version.dataset_id != dataset.id or version.status != 'ready':
            raise HTTPException(status_code=409, detail={'code': 'DATASET_VERSION_UNAVAILABLE', 'message': '数据版本不可用'})
        if not re.fullmatch(rf'dataset_{dataset.id}(?:_v_[a-f0-9]{{32}})?',version.projection_table):
            raise HTTPException(status_code=409, detail={'code': 'DATASET_VERSION_UNAVAILABLE', 'message': '数据版本投影不匹配'})
        return version

    def load_frame(self, dataset: Dataset, columns: list[DatasetColumn], version_id: int | None = None) -> pd.DataFrame:
        version = None
        if self.db is not None:
            version = self.get_version(dataset, version_id)
        elif version_id is not None:
            raise ValueError('pinned versions require an authorized Session')
        name = version.projection_table if version else getattr(dataset, 'projection_table', None) or f'dataset_{int(dataset.id)}'
        if not re.fullmatch(rf'dataset_{int(dataset.id)}(?:_v_[a-f0-9]{{32}})?', name):
            raise ValueError('invalid server projection name')
        if not inspect(self.projection_bind).has_table(name):
            raise LookupError('dataset projection is unavailable')
        table = Table(name, MetaData(), autoload_with=self.projection_bind)
        if version:
            from types import SimpleNamespace
            from app.datasets.schemas import DatasetSchema,column_storage_type
            schema = DatasetSchema.model_validate(version.schema_json)
            columns = [SimpleNamespace(name=column.name,data_type=column_storage_type(column)) for column in schema.columns]
        chunks = []
        with self.projection_bind.connect() as connection:
            for chunk in pd.read_sql(select(table), connection, chunksize=5000, coerce_float=False, dtype_backend='numpy_nullable'):
                chunks.append(chunk)
                if sum(int(part.memory_usage(deep=True).sum()) for part in chunks) > get_settings().dataframe_max_bytes:
                    raise ValueError('dataset exceeds memory budget')
        frame = pd.concat(chunks, ignore_index=True) if chunks else pd.DataFrame(columns=[column.name for column in columns])
        for column in columns:
            if column.data_type in {'date', 'datetime'}:
                frame[column.name] = pd.to_datetime(frame[column.name], errors='raise')
            elif column.data_type == 'integer':
                frame[column.name] = pd.to_numeric(frame[column.name], errors='raise').astype('Int64')
            elif column.data_type == 'decimal':
                frame[column.name] = pd.to_numeric(frame[column.name], errors='raise')
            elif column.data_type == 'boolean':
                frame[column.name] = frame[column.name].astype('boolean')
        frame.attrs['quality_warnings'] = (version.profile_json.get('warnings',[]) if version else getattr(dataset, 'quality_warnings_json', None)) or []
        if version: frame.attrs['original_columns']=[column.original_name for column in schema.columns]
        return frame
