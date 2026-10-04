from __future__ import annotations

import base64
import csv
import html
import io
import json
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Protocol

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.drawing.image import Image as SpreadsheetImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table as PdfTable, TableStyle
from reportlab.platypus.tableofcontents import TableOfContents

from app.reports.schemas import DocumentSection, ReportDocument


MAX_SECTION_ROWS = 100
MAX_EXPORT_ROWS = 50_000
MAX_EXPORT_BYTES = 64 * 1024 * 1024
MAX_IMAGES = 30


def _register_cjk_font() -> str:
    font_name = "DataLensCJK"
    try:
        pdfmetrics.getFont(font_name)
        return font_name
    except KeyError:
        pass
    candidates = [
        Path("C:/Windows/Fonts/msyh.ttc"), Path("C:/Windows/Fonts/msyh.ttf"),
        Path("/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
        Path("/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc"),
    ]
    for candidate in candidates:
        if not candidate.is_file():
            continue
        try:
            pdfmetrics.registerFont(TTFont(font_name, str(candidate), subfontIndex=0))
            pdfmetrics.registerFontFamily(font_name, normal=font_name, bold=font_name,
                                          italic=font_name, boldItalic=font_name)
            return font_name
        except Exception:
            continue
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light", isVertical=False))
    pdfmetrics.registerFontFamily(font_name, normal="STSong-Light", bold="STSong-Light",
                                  italic="STSong-Light", boldItalic="STSong-Light")
    return "STSong-Light"


@dataclass(frozen=True)
class ExportFile:
    file_name: str
    mime_type: str
    content: bytes


@dataclass
class ExportOptions:
    include_raw_data: bool = False
    raw_data: list[dict[str, Any]] = field(default_factory=list)
    chart_files: dict[int, bytes] = field(default_factory=dict)
    max_export_rows: int = MAX_EXPORT_ROWS
    cleaned_data: list[dict[str, Any]] = field(default_factory=list)
    input_snapshots: list[dict[str, Any]] = field(default_factory=list)
    source_plan: dict = field(default_factory=dict)


class ReportExporter(Protocol):
    format: str

    def export(self, document: ReportDocument, options: ExportOptions) -> list[ExportFile]: ...


def _base_name(document: ReportDocument) -> str:
    value = unicodedata.normalize("NFC", document.title).strip()
    value = re.sub(r'[\\/:*?"<>|\x00-\x1f\x7f]', "_", value).strip(" .")[:120]
    return value or "数据分析报告"


def _display(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, allow_nan=False, default=str)
    return str(value)


def _spreadsheet_safe(value: Any) -> Any:
    from decimal import Decimal
    if isinstance(value,(dict,list,tuple,Decimal)): value = _display(value)
    if isinstance(value, str) and value.lstrip(" \t\r\n").startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _tables(section: DocumentSection) -> list[dict[str, Any]]:
    data = section.data or {}
    candidates = data.get("tables")
    if isinstance(candidates, list):
        return [table for table in candidates if isinstance(table, dict)][:20]
    if isinstance(data.get("columns"), list) and isinstance(data.get("rows"), list):
        return [data]
    return []


def _table_shape(table: dict[str, Any], max_rows: int = MAX_SECTION_ROWS) -> tuple[list[str], list[list[Any]]]:
    columns_raw = table.get("columns") or []
    columns = [str(column.get("name", column.get("field", ""))) if isinstance(column, dict)
               else str(column) for column in columns_raw]
    rows_raw = table.get("rows") or []
    rows = [[_spreadsheet_safe(row.get(column)) for column in columns]
            for row in rows_raw[:max_rows] if isinstance(row, dict)]
    return columns, rows


def _docx_set_cell_text(cell, value: Any) -> None:
    cell.text = _display(value)
    for paragraph in cell.paragraphs:
        paragraph.alignment = TA_LEFT
        for run in paragraph.runs:
            run.font.name = "Microsoft YaHei"
            run.font.size = None


class PdfExporter:
    format = "pdf"

    def export(self, document: ReportDocument, options: ExportOptions) -> list[ExportFile]:
        font_name = _register_cjk_font()

        buffer = io.BytesIO()
        page_size = landscape(A4) if any(
            len(_table_shape(table)[0]) > 6
            for section in document.sections for table in _tables(section)
        ) else A4
        class ReportPdf(SimpleDocTemplate):
            def afterFlowable(self, flowable):
                if isinstance(flowable, Paragraph) and hasattr(flowable, '_toc_key'):
                    self.canv.bookmarkPage(flowable._toc_key)
                    self.notify('TOCEntry',(0, flowable.getPlainText(), self.page, flowable._toc_key))
        pdf = ReportPdf(buffer, pagesize=page_size, rightMargin=19 * mm, leftMargin=19 * mm,
                                topMargin=20 * mm, bottomMargin=19 * mm, title=document.title,
                                author=document.author or "DataLens Agent", pageCompression=1)
        source_styles = getSampleStyleSheet()
        styles = {
            "title": ParagraphStyle("DLTitle", parent=source_styles["Title"], fontName=font_name,
                                     fontSize=21, leading=29, textColor=colors.HexColor("#12213a"), alignment=TA_LEFT),
            "subtitle": ParagraphStyle("DLSubtitle", parent=source_styles["Normal"], fontName=font_name,
                                        fontSize=10, leading=16, textColor=colors.HexColor("#64748b")),
            "heading": ParagraphStyle("DLHeading", parent=source_styles["Heading2"], fontName=font_name,
                                       fontSize=14, leading=20, textColor=colors.HexColor("#2f6bff"), spaceBefore=9),
            "body": ParagraphStyle("DLBody", parent=source_styles["BodyText"], fontName=font_name,
                                    fontSize=9.5, leading=15, wordWrap='CJK', textColor=colors.HexColor("#243247")),
            "small": ParagraphStyle("DLSmall", parent=source_styles["BodyText"], fontName=font_name,
                                     fontSize=7.5, leading=10, wordWrap='CJK', textColor=colors.HexColor("#526174")),
        }
        story: list[Any] = [Paragraph(html.escape(document.title), styles["title"]), Spacer(1, 3 * mm)]
        if document.subtitle:
            story += [Paragraph(html.escape(document.subtitle), styles["subtitle"]), Spacer(1, 7 * mm)]
        story.extend([Spacer(1,20*mm),Paragraph('范围：固定输入版本的已保存分析结果。',styles['body']),
            Paragraph(html.escape(document.generated_at.isoformat()),styles['small']),
            Paragraph('报告版本 '+str(document.report_version),styles['small'])])
        if document.metadata.get('user_edited'):
            story.append(Paragraph('此版本包含用户编辑文案；计算表格与证据保留原始来源。',styles['body']))
        if len(document.sections)>2: story.append(PageBreak())
        if document.include_toc and len(document.sections) > 2:
            story.append(Paragraph("目录", styles["heading"]))
            toc = TableOfContents()
            toc.levelStyles = [ParagraphStyle('DLToc',parent=styles['body'],spaceBefore=5,leftIndent=0,firstLineIndent=0)]
            story.append(toc)
            story.append(PageBreak())
        for section in document.sections:
            heading = Paragraph(html.escape(section.title), styles['heading'])
            heading._toc_key = section.section_id
            story.append(heading)
            if section.narrative:
                story.append(Paragraph(html.escape(section.narrative).replace("\n", "<br/>"), styles["body"]))
                story.append(Spacer(1, 3 * mm))
            for table in _tables(section):
                columns, rows = _table_shape(table)
                if not columns:
                    continue
                cells = [[Paragraph(html.escape(column), styles["small"]) for column in columns]]
                cells.extend([[Paragraph(html.escape(_display(value)), styles["small"]) for value in row]
                              for row in rows])
                report_table = PdfTable(cells, colWidths=[pdf.width/len(columns)]*len(columns), repeatRows=1, hAlign="LEFT", splitByRow=1, splitInRow=1)
                report_table.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eaf0fa")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#12213a")),
                    ("FONTNAME", (0, 0), (-1, -1), font_name),
                    ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#d9e1ed")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                    ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]))
                story.extend([report_table, Spacer(1, 4 * mm)])
            for artifact_id in section.artifact_refs[:MAX_IMAGES]:
                image = options.chart_files.get(artifact_id)
                if image:
                    rendered = Image(io.BytesIO(image), width=168 * mm, height=95 * mm, kind="proportional")
                    story.extend([rendered, Spacer(1, 4 * mm)])
        def draw_footer(canvas, doc):
            canvas.saveState()
            canvas.setFont(font_name, 8)
            canvas.setFillColor(colors.HexColor("#65748a"))
            canvas.drawString(19 * mm, 10 * mm, f"DataLens Agent · {document.report_version}")
            canvas.drawString(19 * mm, page_size[1]-12*mm, document.title[:40])
            canvas.drawRightString(page_size[0] - 19 * mm, 10 * mm, str(doc.page))
            canvas.restoreState()
        pdf.multiBuild(story, onFirstPage=draw_footer, onLaterPages=draw_footer)
        return [ExportFile(f"{_base_name(document)}.pdf", "application/pdf", buffer.getvalue())]


class DocxExporter:
    format = "docx"

    def export(self, document: ReportDocument, options: ExportOptions) -> list[ExportFile]:
        from docx import Document
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor

        word = Document()
        section = word.sections[0]
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        normal = word.styles["Normal"]
        normal.font.name = "Microsoft YaHei"
        normal.font.size = Pt(10)
        normal.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
        title = word.add_heading(document.title, 0)
        title.style.font.name = "Microsoft YaHei"
        title.style.font.color.rgb = RGBColor(0, 0, 0)
        if document.subtitle:
            subtitle = word.add_paragraph(document.subtitle)
            subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
        header = section.header.paragraphs[0]
        header.text = "DataLens Agent · 分析报告"
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer.add_run(f"版本 {document.report_version}    ")
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        footer._p.append(field)
        word.add_paragraph('范围：固定输入版本的已保存分析结果。')
        word.add_paragraph(f'生成时间 {document.generated_at.isoformat()}')
        if document.metadata.get('user_edited'): word.add_paragraph('此版本包含用户编辑文案；计算表格与证据保留原始来源。')
        if len(document.sections)>2: word.add_page_break()
        if document.include_toc:
            toc_title = word.add_paragraph('目录')
            toc_title.runs[0].bold = True
            toc_title.runs[0].font.size = Pt(16)
            toc = word.add_paragraph()
            field_toc = OxmlElement('w:fldSimple'); field_toc.set(qn('w:instr'),'TOC \\o "1-2" \\h \\z \\u')
            toc._p.append(field_toc)
            settings_update = OxmlElement('w:updateFields'); settings_update.set(qn('w:val'),'true')
            word.settings.element.append(settings_update)
            word.add_paragraph('打开文档后可更新目录页码。')
            word.add_page_break()
        for block in document.sections:
            word.add_heading(block.title, level=1)
            if block.narrative:
                word.add_paragraph(block.narrative)
            for table_data in _tables(block):
                columns, rows = _table_shape(table_data)
                if not columns:
                    continue
                table = word.add_table(rows=1, cols=len(columns))
                table.style = "Light Shading Accent 1"
                repeat = OxmlElement('w:tblHeader'); repeat.set(qn('w:val'),'true')
                table.rows[0]._tr.get_or_add_trPr().append(repeat)
                for cell, column in zip(table.rows[0].cells, columns):
                    _docx_set_cell_text(cell, column)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    for run in cell.paragraphs[0].runs:
                        run.bold = True
                for row in rows:
                    for cell, value in zip(table.add_row().cells, row):
                        _docx_set_cell_text(cell, value)
                        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for artifact_id in block.artifact_refs[:MAX_IMAGES]:
                image = options.chart_files.get(artifact_id)
                if image:
                    word.add_picture(io.BytesIO(image), width=Inches(6.1))
        buffer = io.BytesIO()
        word.save(buffer)
        return [ExportFile(f"{_base_name(document)}.docx",
                           "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                           buffer.getvalue())]


class ExcelExporter:
    format = "xlsx"

    def export(self, document: ReportDocument, options: ExportOptions) -> list[ExportFile]:
        limit = max(1,min(options.max_export_rows,MAX_EXPORT_ROWS))
        if len(options.raw_data) > limit or len(options.cleaned_data) > limit:
            from dataclasses import replace
            count = max((len(options.raw_data)+limit-1)//limit,(len(options.cleaned_data)+limit-1)//limit)
            outputs, parts = [], []
            for index in range(count):
                selected = replace(options,raw_data=options.raw_data[index*limit:(index+1)*limit],
                    cleaned_data=options.cleaned_data[index*limit:(index+1)*limit])
                file = self.export(document,selected)[0]
                name = f'{_base_name(document)}_part_{index+1}.xlsx'
                outputs.append(ExportFile(name,file.mime_type,file.content))
                parts.append({'file_name':name,'row_count':len(selected.raw_data),'cleaned_row_count':len(selected.cleaned_data),'raw_offset':index*limit})
            outputs.append(ExportFile(f'{_base_name(document)}_manifest.json','application/json',json.dumps({
                'dataset_versions':document.metadata.get('dataset_versions',[]),'parts':parts},ensure_ascii=False).encode()))
            return outputs
        workbook = Workbook()
        summary = workbook.active
        summary.title = "Summary"
        summary.append(["报告", _spreadsheet_safe(document.title)])
        summary.append(["数据集", _spreadsheet_safe(document.metadata.get("dataset_name", ""))])
        summary.append(["数据集版本", document.metadata.get("dataset_version_id")])
        summary.append(["生成时间", document.generated_at.isoformat()])
        summary.append(["报告版本", document.report_version])
        if document.subtitle:
            summary.append(["说明", _spreadsheet_safe(document.subtitle)])
        analysis = workbook.create_sheet("Analysis")
        metrics = workbook.create_sheet("Metrics")
        charts = workbook.create_sheet("Charts")
        kpi = workbook.create_sheet('KPI')
        quality = workbook.create_sheet('Data Quality')
        cleaned = workbook.create_sheet('Cleaned Data')
        insights = workbook.create_sheet('Insights')
        quality.append(['输入','固定版本','质量概要','清洗操作'])
        for item in options.input_snapshots:
            quality.append([_spreadsheet_safe(item.get('alias','primary')),item.get('dataset_version_id'),
                _spreadsheet_safe(_display(item.get('profile',{}))),_spreadsheet_safe(_display(item.get('transformations',[])))])
        insights.append(['章节','结论或限制','证据'])
        for section in document.sections:
            if section.content_type == 'insights': insights.append([section.title,_spreadsheet_safe(section.narrative or ''),', '.join(section.evidence_ids)])
        native_chart_source = None
        for section in document.sections:
            analysis.append([_spreadsheet_safe(section.title)])
            if section.narrative:
                analysis.append([_spreadsheet_safe(section.narrative)])
            if section.evidence_ids:
                analysis.append(["Evidence", ", ".join(section.evidence_ids)])
            for data_table in _tables(section):
                columns, rows = _table_shape(data_table)
                if not columns:
                    continue
                start = 1 if metrics.max_row == 1 and metrics.cell(1, 1).value is None else metrics.max_row + 1
                for column_index, column in enumerate(columns, start=1):
                    metrics.cell(row=start, column=column_index, value=_spreadsheet_safe(column))
                for row in rows:
                    metrics.append(row)
                end = metrics.max_row
                if end > start:
                    ref = f"A{start}:{get_column_letter(len(columns))}{end}"
                    table = Table(displayName=f"MetricsTable{metrics.max_row}", ref=ref)
                    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
                    metrics.add_table(table)
                    if native_chart_source is None and len(columns) >= 2 and all(
                        isinstance(value, (int, float)) and not isinstance(value, bool)
                        for row in rows for value in row[1:]):
                        native_chart_source = (start, end, len(columns))
        analysis.column_dimensions["A"].width = 100
        for column_index in range(1, metrics.max_column + 1):
            metrics.column_dimensions[get_column_letter(column_index)].width = 24
        if options.include_raw_data:
            raw = workbook.create_sheet("Raw Data")
            if options.raw_data:
                headers = list(dict.fromkeys(key for row in options.raw_data for key in row))
                raw.append([_spreadsheet_safe(header) for header in headers])
                for record in options.raw_data:
                    raw.append([_spreadsheet_safe(record.get(header)) for header in headers])
                self._style_sheet(raw)
        if options.cleaned_data:
            headers = list(dict.fromkeys(key for row in options.cleaned_data for key in row))
            cleaned.append([_spreadsheet_safe(header) for header in headers])
            for row in options.cleaned_data: cleaned.append([_spreadsheet_safe(row.get(key)) for key in headers])
        else:
            cleaned.append(['状态']); cleaned.append(['当前输入未采用清洗版本；原始数据见 Raw Data。'])
        for row in metrics.values: kpi.append(list(row))
        if native_chart_source:
            start, end, column_count = native_chart_source
            native = BarChart()
            native.type = "col"
            native.style = 10
            native.title = document.title[:80]
            native.y_axis.title = "数值"
            native.x_axis.title = "分类"
            native.height, native.width = 8, 16
            native_data = Reference(metrics, min_col=2, max_col=column_count,
                                    min_row=start, max_row=end)
            native_categories = Reference(metrics, min_col=1, min_row=start + 1, max_row=end)
            native.add_data(native_data, titles_from_data=True)
            native.set_categories(native_categories)
            charts.add_chart(native, "A2")
        image_anchor_row = 20 if native_chart_source else 2
        for artifact_id, image_content in options.chart_files.items():
            if type(artifact_id) is not int or artifact_id not in document.artifacts:
                continue
            picture = SpreadsheetImage(io.BytesIO(image_content))
            picture.width, picture.height = 640, 384
            charts.add_image(picture, f"A{image_anchor_row}")
            image_anchor_row += 22
        charts.column_dimensions["A"].width = 22
        for sheet in (summary, analysis, metrics, charts, kpi, quality, cleaned, insights):
            self._style_sheet(sheet)
        buffer = io.BytesIO()
        workbook.save(buffer)
        return [ExportFile(f"{_base_name(document)}.xlsx",
                           "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                           buffer.getvalue())]

    @staticmethod
    def _style_sheet(sheet) -> None:
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2F6BFF")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for row in sheet.iter_rows():
            for cell in row:
                if cell.row > 1:
                    cell.font = Font(name="Microsoft YaHei", color="12213A")
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
        for index in range(1, sheet.max_column + 1):
            letter = get_column_letter(index)
            if not sheet.column_dimensions[letter].width:
                sheet.column_dimensions[letter].width = min(36, max(12, max(
                    (len(_display(cell.value)) for cell in sheet[letter] if cell.value is not None), default=12) + 2))
        if sheet.max_row > 1 and sheet.title in {"Metrics", "Raw Data"}:
            sheet.auto_filter.ref = sheet.dimensions


class HtmlExporter:
    format = "html"

    def export(self, document: ReportDocument, options: ExportOptions) -> list[ExportFile]:
        parts = ["<!doctype html><html lang=\"zh-CN\"><head><meta charset=\"utf-8\">",
                 "<meta name=\"viewport\" content=\"width=device-width, initial-scale=1\"><title>",
                 html.escape(document.title), "</title><style>",
                 "body{margin:0;background:#f3f6fa;color:#12213a;font:16px/1.65 'Microsoft YaHei',system-ui,sans-serif}"
                 "main{max-width:980px;margin:40px auto;padding:48px;background:white;border:1px solid #d9e1ed}"
                 "h1{font-size:30px}h2{margin-top:36px;color:#2457cd;font-size:20px}"
                 "table{border-collapse:collapse;width:100%;margin:18px 0;font-size:14px}"
                 "th,td{border:1px solid #d9e1ed;padding:8px;text-align:left}th{background:#eaf0fa}"
                 "img{display:block;max-width:100%;height:auto;margin:20px auto}aside{color:#5b687b;font-size:13px}"
                 "@media print{body{background:white}main{margin:0;max-width:none;border:0;padding:0}}"
                 "</style></head><body><main><h1>", html.escape(document.title), "</h1>"]
        if document.subtitle:
            parts.extend(["<p>", html.escape(document.subtitle), "</p>"])
        if document.include_toc and document.sections:
            parts.append("<nav aria-label=\"目录\"><strong>目录</strong><ol>")
            for section in document.sections:
                parts.extend(["<li><a href=\"#", html.escape(section.section_id), "\">",
                              html.escape(section.title), "</a></li>"])
            parts.append("</ol></nav>")
        for section in document.sections:
            parts.extend(["<section id=\"", html.escape(section.section_id), "\"><h2>",
                          html.escape(section.title), "</h2>"])
            if section.narrative:
                parts.extend(["<p>", html.escape(section.narrative).replace("\n", "<br>"), "</p>"])
            for table_data in _tables(section):
                columns, rows = _table_shape(table_data)
                if columns:
                    parts.append("<table><thead><tr>")
                    parts.extend("<th>" + html.escape(column) + "</th>" for column in columns)
                    parts.append("</tr></thead><tbody>")
                    for row in rows:
                        parts.append("<tr>")
                        parts.extend("<td>" + html.escape(_display(value)) + "</td>" for value in row)
                        parts.append("</tr>")
                    parts.append("</tbody></table>")
            for artifact_id in section.artifact_refs[:MAX_IMAGES]:
                image = options.chart_files.get(artifact_id)
                if image:
                    encoded = base64.b64encode(image).decode("ascii")
                    parts.extend(["<img alt=\"分析图表\" src=\"data:image/png;base64,", encoded, "\">"])
            if section.evidence_ids:
                parts.extend(["<aside>数据依据：", html.escape(" · ".join(section.evidence_ids)), "</aside>"])
            parts.append("</section>")
        parts.extend(["<footer><aside>报告版本 ", str(document.report_version), "</aside></footer></main></body></html>"])
        content = "".join(parts).encode("utf-8")
        return [ExportFile(f"{_base_name(document)}.html", "text/html; charset=utf-8", content)]


class MarkdownExporter:
    format = "markdown"

    def export(self, document: ReportDocument, options: ExportOptions) -> list[ExportFile]:
        lines = [f"# {document.title}", ""]
        if document.subtitle:
            lines += [document.subtitle, ""]
        for section in document.sections:
            lines += [f"## {section.title}", ""]
            if section.narrative:
                lines += [section.narrative, ""]
            for table_data in _tables(section):
                columns, rows = _table_shape(table_data)
                if not columns:
                    continue
                lines.append("| " + " | ".join(columns) + " |")
                lines.append("| " + " | ".join("---" for _ in columns) + " |")
                lines.extend("| " + " | ".join(_display(value).replace("|", "\\|") for value in row) + " |"
                             for row in rows)
                lines.append("")
            if section.evidence_ids:
                lines += ["数据依据：" + "、".join(section.evidence_ids), ""]
            for artifact_id in section.artifact_refs[:MAX_IMAGES]:
                if artifact_id in options.chart_files:
                    lines += [f"![图表](assets/{artifact_id}.png)", ""]
        markdown = "\n".join(lines).encode("utf-8")
        outputs = [ExportFile(f"{_base_name(document)}.md", "text/markdown; charset=utf-8", markdown)]
        if options.chart_files:
            archive = io.BytesIO()
            with zipfile.ZipFile(archive, "w", compression=zipfile.ZIP_DEFLATED) as bundle:
                bundle.writestr(f"{_base_name(document)}.md", markdown)
                for artifact_id, image in options.chart_files.items():
                    if type(artifact_id) is int and artifact_id > 0 and isinstance(image, bytes):
                        bundle.writestr(f"assets/{artifact_id}.png", image)
            outputs.append(ExportFile(f"{_base_name(document)}_完整包.zip", "application/zip", archive.getvalue()))
        return outputs


class CsvExporter:
    format = "csv"

    def export(self, document: ReportDocument, options: ExportOptions) -> list[ExportFile]:
        outputs: list[ExportFile] = []
        for section in document.sections:
            for index, data_table in enumerate(_tables(section), start=1):
                columns, rows = _table_shape(data_table, max_rows=min(MAX_SECTION_ROWS, options.max_export_rows))
                if not columns:
                    continue
                buffer = io.StringIO(newline="")
                writer = csv.writer(buffer, lineterminator="\r\n")
                writer.writerow(columns)
                writer.writerows(rows)
                base = _base_name(document)
                suffix = f"_{index}" if len(_tables(section)) > 1 else ""
                name = f"{base}_{section.section_id}{suffix}.csv"
                outputs.append(ExportFile(name, "text/csv; charset=utf-8", b"\xef\xbb\xbf" + buffer.getvalue().encode("utf-8")))
        if not outputs:
            raise ValueError("CSV_TABLE_NOT_FOUND")
        return outputs


class JsonExporter:
    format = "json"

    def export(self, document: ReportDocument, options: ExportOptions) -> list[ExportFile]:
        content = json.dumps(document.model_dump(mode="json"), ensure_ascii=False,
                             allow_nan=False, separators=(",", ":")).encode("utf-8")
        return [ExportFile(f"{_base_name(document)}.json", "application/json", content)]


class ExporterRegistry:
    def __init__(self, exporters: list[ReportExporter] | None = None):
        self._exporters: dict[str, ReportExporter] = {}
        from app.reports.code_exporter import CodeExporter
        for exporter in exporters or [PdfExporter(), DocxExporter(), ExcelExporter(), HtmlExporter(),
                                      MarkdownExporter(), CsvExporter(), JsonExporter(), CodeExporter('python'), CodeExporter('sql')]:
            if exporter.format in self._exporters:
                raise ValueError("Duplicate report exporter")
            self._exporters[exporter.format] = exporter

    def export(self, format: str, document: ReportDocument,
               options: ExportOptions | None = None) -> list[ExportFile]:
        try:
            exporter = self._exporters[format.lower().strip()]
        except (KeyError, AttributeError):
            raise ValueError("REPORT_FORMAT_UNSUPPORTED") from None
        files = exporter.export(document, options or ExportOptions())
        if not files or any(not item.content or len(item.content) > MAX_EXPORT_BYTES for item in files):
            raise ValueError("REPORT_EXPORT_SIZE_EXCEEDED")
        return files

    @property
    def formats(self) -> tuple[str, ...]:
        return tuple(self._exporters)


def exporter_registry() -> ExporterRegistry:
    return ExporterRegistry()
