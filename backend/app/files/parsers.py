from pathlib import Path, PurePosixPath
import zipfile

MAX_TEXT = 2_000_000
MAX_PAGES = 200
MAX_CANDIDATES = 100
MAX_CELLS = 100_000
MAX_CELL_TEXT = 10_000


class FileParseError(Exception):
    def __init__(self, code, message):
        super().__init__(message)
        self.code, self.message = code, message


def fail(code, message):
    raise FileParseError(code, message)


def validate_signature(path, kind):
    with path.open('rb') as source:
        head = source.read(8)
        text_has_nul = b'\x00' in head
        if kind in {'txt','csv','tsv','json','jsonl'}:
            while chunk := source.read(1024*1024):
                if b'\x00' in chunk:
                    text_has_nul = True
                    break
    signatures = {'pdf': b'%PDF-', 'docx': b'PK\x03\x04', 'xlsx': b'PK\x03\x04',
                  'parquet': b'PAR1', 'xls': bytes.fromhex('D0CF11E0A1B11AE1')}
    if kind in signatures and not head.startswith(signatures[kind]):
        fail('FILE_SIGNATURE_INVALID', '文件签名与扩展名不一致')
    if kind in {'txt','csv','tsv','json','jsonl'} and (text_has_nul or head.startswith((b'MZ',b'PK',b'%PDF-'))):
        fail('FILE_SIGNATURE_INVALID', '文本文件包含不支持的二进制内容')


def check_docx(path):
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            expanded = sum(e.file_size for e in entries)
            if expanded > 100*1024*1024 or expanded/max(1,sum(e.compress_size for e in entries)) > 100:
                fail('FILE_ARCHIVE_LIMIT', '文档解压体积或压缩比超过限制')
            for entry in entries:
                name = entry.filename.replace('\\','/')
                if PurePosixPath(name).is_absolute() or '..' in PurePosixPath(name).parts or ':' in name:
                    fail('FILE_ARCHIVE_PATH', '文档包含不安全的归档路径')
                if name.lower().endswith('vbaproject.bin'):
                    fail('FILE_MACRO_UNSUPPORTED', '不支持含宏文档')
            if 'word/document.xml' not in archive.namelist():
                fail('FILE_SIGNATURE_INVALID', '文件不是 DOCX 文档')
    except zipfile.BadZipFile:
        fail('FILE_PARSE_FAILED', '文档归档损坏')


def parse_document(path: Path, kind: str):
    validate_signature(path, kind)
    candidates, parts, cells = [], [], 0
    def add_text(value):
        parts.append(value)
        if sum(len(p) for p in parts) > MAX_TEXT:
            fail('FILE_TEXT_LIMIT', '文档文本超过解析限制')
    def add_table(rows, location):
        nonlocal cells
        if not rows: return
        rows = [[str(c or '') for c in row] for row in rows]
        width=max(len(row) for row in rows)
        # The budget covers the rectangular candidate we retain, including
        # padding for omitted DOCX grid cells, rather than only source cells.
        cells += len(rows)*width
        if len(candidates) >= MAX_CANDIDATES or cells > MAX_CELLS or any(len(c)>MAX_CELL_TEXT for row in rows for c in row):
            fail('FILE_TABLE_LIMIT', '表格候选或单元格超过解析限制')
        if width>200 or len(rows)>100001:
            fail('FILE_TABLE_LIMIT', '候选表格超过行列限制')
        rows=[row+['']*(width-len(row)) for row in rows]
        candidates.append({'id':str(len(candidates)+1),'rows':rows,'location':location,'requires_header_confirmation':True})
    try:
        if kind=='txt':
            try: add_text(path.read_text(encoding='utf-8-sig'))
            except UnicodeError: fail('FILE_ENCODING_UNSUPPORTED','TXT 请另存为 UTF-8')
            if '\x00' in parts[0]: fail('FILE_SIGNATURE_INVALID','TXT 包含二进制空字符')
        elif kind=='docx':
            check_docx(path)
            from docx import Document
            doc=Document(path)
            for paragraph in doc.paragraphs: add_text(paragraph.text)
            for index, table in enumerate(doc.tables):
                add_table([[cell.text for cell in row.cells] for row in table.rows], {'table_index':index})
                add_text('\n'.join('\t'.join(cell.text for cell in row.cells) for row in table.rows))
        elif kind=='pdf':
            from pypdf import PdfReader
            reader=PdfReader(path)
            if reader.is_encrypted: fail('FILE_ENCRYPTED_UNSUPPORTED','不支持加密 PDF')
            if len(reader.pages)>MAX_PAGES: fail('FILE_PAGE_LIMIT','PDF 超过 200 页解析限制')
            for page in reader.pages: add_text(page.extract_text() or '')
            if not ''.join(parts).strip(): fail('FILE_OCR_UNSUPPORTED','PDF 没有可提取文本；暂不支持扫描件 OCR')
            import pdfplumber
            with pdfplumber.open(path) as pdf:
                for index,page in enumerate(pdf.pages):
                    for table in page.find_tables(): add_table(table.extract(), {'page':index+1,'bbox':list(table.bbox)})
        else: fail('FILE_TYPE_UNSUPPORTED','不支持的文档格式')
    except FileParseError: raise
    except Exception: fail('FILE_PARSE_FAILED','文档无法解析，请检查文件格式')
    text='\n'.join(parts)
    return {'text':text,'candidates':candidates}
