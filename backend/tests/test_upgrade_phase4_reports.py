import ast
import io
import json
import zipfile
from datetime import UTC, datetime
import pytest
from openpyxl import load_workbook
from docx import Document
from pypdf import PdfReader
from test_analysis_api import analysis_context
from test_upgrade_phase4_api import setup_source
from app.reports.schemas import ReportSpec
from app.reports.builder import ReportBuilder
from app.reports.exporters import ExportOptions, exporter_registry
from app.reports.service import create_report, export_report_version
from app.models import AnalysisRecord


def report_source():
    return {'record_id':1,'dataset_id':1,'dataset_version_id':1,'status':'succeeded',
        'answer':'收入合计为 12。','report':{'tables':[{'columns':['收入'],'rows':[{'收入':12}]}]},
        'evidence':[{'evidence_id':'1:kpi:revenue','dataset_version_id':1,'value':'12'}],
        'artifacts':[], 'dataset_versions':[{'alias':'primary','dataset_id':1,'dataset_version_id':1}]}


def test_templates_have_distinct_relevant_sections_and_evidence_appendix():
    documents = [ReportBuilder().build(ReportSpec(title='分析',dataset_id=1,dataset_version_id=1,template=name),[report_source()])
        for name in ['quick','detailed','executive','technical','data_quality','forecast']]
    assert len({tuple(s.section_id for s in d.sections) for d in documents}) == 6
    assert all(any(s.section_id=='evidence_appendix' for s in d.sections) for d in documents)
    assert any(s.section_id=='forecast_validation' for s in documents[-1].sections)


def test_service_xlsx_uses_fixed_input_not_analysis_summary(analysis_context):
    client,sessions,_ = analysis_context
    rid,sid,aid,did,vid = setup_source(client,sessions)
    with sessions() as db:
        record = db.get(AnalysisRecord,rid)
        report,version = create_report(db,record.user_id,sid,ReportSpec(title='完整明细',dataset_id=did,dataset_version_id=vid),[rid])
        files = export_report_version(db,report,version,'xlsx')
        from app.artifacts.storage import LocalArtifactStorage
        from app.config import get_settings
        content = LocalArtifactStorage(get_settings().artifact_dir).read(files[0].storage_key)
        book = load_workbook(io.BytesIO(content))
        assert {'Summary','Raw Data','Cleaned Data','Data Quality','KPI','Analysis','Charts','Insights'} <= set(book.sheetnames)
        assert book['Raw Data'].max_row == 3
        assert [r[0] for r in list(book['Raw Data'].values)[1:]] == [1,2]
        assert files[0].expires_at is None
        assert version.dataset_versions_json[0]['dataset_version_id'] == vid


def test_raw_rows_split_with_manifest_and_formula_safety():
    doc = ReportBuilder().build(ReportSpec(title='明细',dataset_id=1,dataset_version_id=1),[report_source()])
    files = exporter_registry().export('xlsx',doc,ExportOptions(include_raw_data=True,
        raw_data=[{'name':'=cmd','value':x} for x in range(5)], max_export_rows=2))
    assert len(files) == 4
    manifest = json.loads(next(f.content for f in files if f.file_name.endswith('.json')))
    assert sum(part['row_count'] for part in manifest['parts']) == 5
    for file in files:
        if file.file_name.endswith('.xlsx'):
            sheet = load_workbook(io.BytesIO(file.content))['Raw Data']
            assert sheet.max_row <= 3 and sheet['A2'].value == "'=cmd"


def test_docx_toc_and_repeat_header_pdf_real_page_numbers_and_evidence():
    doc = ReportBuilder().build(ReportSpec(title='中文多页报告',dataset_id=1,dataset_version_id=1,template='detailed'),[report_source()])
    doc.sections[0].narrative = '可追溯分析证据。\n'*130
    outputs = {fmt:exporter_registry().export(fmt,doc)[0].content for fmt in ['pdf','docx']}
    pdf = PdfReader(io.BytesIO(outputs['pdf']))
    assert len(pdf.pages)>=3
    text = '\n'.join(p.extract_text() for p in pdf.pages)
    assert '证据附录' in text and '收入' in text
    with zipfile.ZipFile(io.BytesIO(outputs['docx'])) as archive:
        xml = archive.read('word/document.xml').decode()
        assert 'TOC' in xml and 'tblHeader' in xml


def test_python_export_is_parameterized_and_sql_omits_credentials(tmp_path):
    from app.reports.code_exporter import export_plan
    plan = {'version':'3.0','inputs':[{'alias':'primary','dataset_id':1,'dataset_version_id':4}],
        'steps':[{'step_id':'kpi','tool_name':'aggregate','status':'COMPLETED','input_alias':'primary',
                  'arguments':{'metrics':[{'column':'amount','aggregation':'sum'}]},'source_ref':'dataset'}]}
    files = export_plan(plan,'python','analysis')
    code = files[0].content.decode()
    ast.parse(code)
    assert '--input' in code and 'api_key' not in code.lower()
    assert any(f.file_name.endswith('.md') for f in files)
    with pytest.raises(ValueError,match='SQL_EXPORT_NO_VERIFIED_SELECT'):
        export_plan(plan,'sql','analysis')
    import os, subprocess, sys
    from pathlib import Path
    schema={'columns':[{'name':'amount','original_name':'amount','dtype':'object','storage_type':'decimal'}]}
    plan['steps'][0].update(tool_name='kpi_analysis',arguments={'metrics':{'revenue':'amount'},'currency':'CNY','unit':'元'})
    plan['semantic_snapshot']=[{'column':'amount','concept':'revenue','currency':'CNY','unit':'元','source':'user','dataset_version_id':4}]
    code=export_plan(plan,'python','exact',[{**plan['inputs'][0],'schema':schema}])[0].content
    script=tmp_path/'replay.py'; script.write_bytes(code)
    csv=tmp_path/'fixed.csv'; csv.write_text('amount\n9007199254740993.12\n0.03\n',encoding='utf-8')
    output=tmp_path/'result.json'
    subprocess.run([sys.executable,str(script),'--input',f'primary={csv}','--output',str(output)],
        env={**os.environ,'PYTHONPATH':str(Path(__file__).resolve().parents[1])},check=True,capture_output=True)
    assert json.loads(output.read_text(encoding='utf-8'))['kpi']['metrics']['revenue']['value']=='9007199254740993.15'


def test_join_lineage_pins_right_input_and_exports_both_originals(analysis_context):
    from app.models import DatasetVersion, Dataset
    from app.artifacts.storage import LocalArtifactStorage
    from app.config import get_settings
    client,sessions,_ = analysis_context
    rid,sid,aid,did,vid = setup_source(client,sessions)
    right_id = client.post('/api/v1/datasets/upload',files={'file':('right.csv',b'x\n7\n8\n9\n','text/csv')}).json()['data']['id']
    with sessions() as db:
        right = db.get(Dataset,right_id); base = db.get(DatasetVersion,vid)
        joined = DatasetVersion(dataset_id=did,version_number=2,parent_version_id=vid,status='ready',
            source_kind='tool_transform',projection_schema=base.projection_schema,projection_table=base.projection_table,
            schema_json=base.schema_json,profile_json=base.profile_json,transformations_json=[{
                'source_versions':[{'dataset_id':did,'version_id':vid},{'dataset_id':right_id,'version_id':right.current_version_id}]}],
            original_available=False,created_at=datetime.now(UTC))
        db.add(joined); db.flush()
        record = db.get(AnalysisRecord,rid); record.dataset_version_id=joined.id
        db.flush()
        report,version = create_report(db,record.user_id,sid,ReportSpec(title='合并来源',dataset_id=did,dataset_version_id=joined.id),[rid])
        assert {item['dataset_id'] for item in version.dataset_versions_json}=={did,right_id}
        file = export_report_version(db,report,version,'xlsx')[0]
        book = load_workbook(io.BytesIO(LocalArtifactStorage(get_settings().artifact_dir).read(file.storage_key)))
        assert book['Raw Data'].max_row==6
        assert book['Cleaned Data'].max_row==3


def test_missing_bound_fact_source_cannot_become_report_prose(analysis_context):
    from app.reports.service import ReportError
    client,sessions,_ = analysis_context
    rid,sid,aid,did,vid = setup_source(client,sessions)
    with sessions() as db:
        record=db.get(AnalysisRecord,rid)
        record.report_json={'findings':[{'kind':'bound_fact','reference':{'step_id':'deleted','path':['value'],'key':'value'}}]}
        with pytest.raises(ReportError,match='REPORT_EVIDENCE_UNAVAILABLE'):
            create_report(db,record.user_id,sid,ReportSpec(title='缺失证据',dataset_id=did,dataset_version_id=vid),[rid])
