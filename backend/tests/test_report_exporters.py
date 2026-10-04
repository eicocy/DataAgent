import csv
import io
import json
import zipfile
from datetime import UTC, datetime
from pathlib import Path

import pytest
import reportlab
from reportlab.lib import fonts as font_mappings
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

from openpyxl import load_workbook
from pypdf import PdfReader
from PIL import Image

from app.reports.exporters import ExportOptions, exporter_registry
from app.reports.schemas import DocumentSection, ReportDocument
from app.reports import exporters


@pytest.fixture
def isolated_font_candidates(monkeypatch, tmp_path):
    # Keep registration and font-family mappings independent of test order.
    monkeypatch.setattr(pdfmetrics, "_fonts", {
        name: font for name, font in pdfmetrics._fonts.items() if name != "DataLensCJK"
    })
    monkeypatch.setattr(pdfmetrics, "_dynFaceNames", {})
    monkeypatch.setattr(font_mappings, "_tt2ps_map", font_mappings._tt2ps_map.copy())
    monkeypatch.setattr(font_mappings, "_ps2tt_map", font_mappings._ps2tt_map.copy())
    monkeypatch.setattr(exporters, "Path", lambda path: tmp_path / Path(path).name)
    return tmp_path, (Path(reportlab.__file__).parent / "fonts" / "Vera.ttf").read_bytes()


def test_pdf_registers_embeddable_linux_font_without_windows_fonts(isolated_font_candidates):
    directory, true_type_bytes = isolated_font_candidates
    (directory / "wqy-zenhei.ttc").write_bytes(true_type_bytes)

    assert exporters._register_cjk_font() == "DataLensCJK"
    assert isinstance(pdfmetrics.getFont("DataLensCJK"), TTFont)


def test_pdf_tries_next_font_when_linux_font_cannot_load(isolated_font_candidates):
    directory, true_type_bytes = isolated_font_candidates
    (directory / "wqy-zenhei.ttc").write_bytes(b"invalid font")
    (directory / "NotoSansCJK-Regular.ttc").write_bytes(true_type_bytes)

    assert exporters._register_cjk_font() == "DataLensCJK"
    assert isinstance(pdfmetrics.getFont("DataLensCJK"), TTFont)


def document():
    return ReportDocument(
        title="2026 年销售报告",
        report_type="sales",
        theme="professional",
        metadata={"dataset_id": 7, "dataset_version_id": 19, "dataset_name": "sales.xlsx"},
        sections=[DocumentSection(
            section_id="summary", title="执行摘要", content_type="insights",
            narrative="华南销售额为 40。",
            data={"evidence_ids": ["run-5:summary:amount"]},
        ), DocumentSection(
            section_id="metrics", title="区域汇总", content_type="metrics",
            data={"tables": [{"columns": ["region", "sales"], "rows": [
                {"region": "华南", "sales": 40}, {"region": "华东", "sales": 60},
            ]}]},
        )],
        evidence=[{"evidence_id": "run-5:summary:amount", "dataset_version_id": 19}],
        report_id=9,
        report_version=2,
        generated_at=datetime(2026, 10, 3, tzinfo=UTC),
    )


def test_exporter_registry_generates_openable_formats_from_one_document():
    registry = exporter_registry()
    report = document()
    expected = {
        "pdf": ("pdf", b"%PDF-"),
        "docx": ("docx", b"PK"),
        "xlsx": ("xlsx", b"PK"),
        "html": ("html", b"<!doctype html"),
        "markdown": ("md", b"# 2026 "),
        "csv": ("csv", b"\xef\xbb\xbf"),
        "json": ("json", b"{"),
    }

    outputs = {name: registry.export(name, report, ExportOptions()) for name in expected}

    for name, (extension, prefix) in expected.items():
        assert outputs[name]
        assert outputs[name][0].file_name.endswith(f".{extension}")
        assert outputs[name][0].content.startswith(prefix)
    report.sections[0].artifact_refs = [42]
    markdown_outputs = registry.export(
        "markdown", report, ExportOptions(chart_files={42: b"PNG fixture"})
    )
    assert len(markdown_outputs) == 2
    with zipfile.ZipFile(io.BytesIO(markdown_outputs[1].content)) as archive:
        assert "2026 年销售报告.md" in archive.namelist()
        assert archive.read("assets/42.png") == b"PNG fixture"
    with zipfile.ZipFile(io.BytesIO(outputs["docx"][0].content)) as archive:
        assert "word/document.xml" in archive.namelist()
    with zipfile.ZipFile(io.BytesIO(outputs["xlsx"][0].content)) as archive:
        assert "xl/workbook.xml" in archive.namelist()


def test_excel_report_contains_summary_analysis_and_metrics_tables():
    exported = exporter_registry().export("xlsx", document(), ExportOptions())[0]
    workbook = load_workbook(io.BytesIO(exported.content), data_only=False)

    assert {"Summary", "Analysis", "Metrics", "Charts"} <= set(workbook.sheetnames)
    summary = workbook["Summary"]
    assert any(cell.value == "2026 年销售报告" for row in summary for cell in row)
    metrics = workbook["Metrics"]
    assert metrics["A1"].value == "region"
    assert metrics["B2"].value == 40
    assert len(workbook["Charts"]._charts) == 1


def test_excel_export_embeds_chart_image_and_preserves_native_metric_chart():
    report = document()
    report.artifacts = [42]
    report.sections[0].artifact_refs = [42]
    image_buffer = io.BytesIO()
    Image.new("RGB", (16, 9), "#2F6BFF").save(image_buffer, format="PNG")
    exported = exporter_registry().export("xlsx", report,
        ExportOptions(chart_files={42: image_buffer.getvalue()}))[0]
    workbook = load_workbook(io.BytesIO(exported.content))

    assert len(workbook["Charts"]._charts) == 1
    assert len(workbook["Charts"]._images) == 1


def test_html_report_escapes_narrative_and_json_keeps_evidence_links():
    report = document()
    report.sections[0].narrative = "<script>alert('x')</script>"

    html = exporter_registry().export("html", report, ExportOptions())[0].content.decode()
    payload = json.loads(exporter_registry().export("json", report, ExportOptions())[0].content)

    assert "<script>alert" not in html
    assert "&lt;script&gt;" in html
    assert payload["evidence"][0]["evidence_id"] == "run-5:summary:amount"
    csv_files = exporter_registry().export("csv", document(), ExportOptions())
    rows = list(csv.DictReader(io.StringIO(csv_files[0].content.decode("utf-8-sig"))))
    assert rows == [{"region": "华南", "sales": "40"}, {"region": "华东", "sales": "60"}]


def test_pdf_contains_chinese_text_and_embedded_true_type_font():
    exported = exporter_registry().export("pdf", document(), ExportOptions())[0]
    reader = PdfReader(io.BytesIO(exported.content))
    page = reader.pages[0]
    fonts = page["/Resources"]["/Font"].get_object().values()
    embedded = []
    for font_ref in fonts:
        font = font_ref.get_object()
        descriptor = font.get("/FontDescriptor")
        if descriptor:
            descriptor = descriptor.get_object()
            embedded.extend(descriptor.get(name) for name in ("/FontFile", "/FontFile2", "/FontFile3") if descriptor.get(name))

    assert "销售报告" in (page.extract_text() or "")
    assert embedded
