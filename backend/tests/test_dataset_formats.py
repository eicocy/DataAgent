import json
from io import BytesIO

import pytest
import pandas as pd
from openpyxl import Workbook
import pyarrow as pa
import pyarrow.parquet as parquet

from app.services.datasets import DatasetParseError, parse_file
from sqlalchemy import select
from test_analysis_api import analysis_context
from app.models import Dataset


def test_tsv_uses_tab_delimiter_and_existing_column_normalization(tmp_path):
    path = tmp_path / "sales.tsv"
    path.write_text("区域\t销售额\n华南\t12\n华东\t15\n", encoding="utf-8")

    frame = parse_file(path, "tsv")

    assert list(frame.columns) == ["column_1", "column_2"]
    assert frame.iloc[0].tolist() == ["华南", 12]
    assert frame.attrs["source_format"] == "tsv"


@pytest.mark.parametrize("payload", [
    '[{"region":"South","sales":12},{"region":"East","sales":15}]',
    '{"region":"South","sales":12}\n{"region":"East","sales":15}\n',
])
def test_json_accepts_flat_records_and_json_lines(tmp_path, payload):
    path = tmp_path / "sales.json"
    path.write_text(payload, encoding="utf-8")

    frame = parse_file(path, "json")

    assert frame.to_dict(orient="records") == [
        {"region": "South", "sales": 12}, {"region": "East", "sales": 15}
    ]
    assert frame.attrs["source_format"] == "json"


def test_json_rejects_nested_values_instead_of_stringifying_them(tmp_path):
    path = tmp_path / "nested.json"
    path.write_text(json.dumps([{"region": {"name": "South"}, "sales": 12}]), encoding="utf-8")

    with pytest.raises(DatasetParseError) as exc:
        parse_file(path, "json")

    assert exc.value.code == "DATASET_NESTED_JSON_UNSUPPORTED"


def test_xlsx_honors_selected_nonempty_worksheet(tmp_path):
    path = tmp_path / "sales.xlsx"
    workbook = Workbook()
    first = workbook.active
    first.title = "Readme"
    first.append(["说明"])
    sheet = workbook.create_sheet("Actual sales")
    sheet.append(["region", "sales"])
    sheet.append(["South", 12])
    workbook.save(path)
    workbook.close()

    frame = parse_file(path, "xlsx", sheet_name="Actual sales")

    assert frame.to_dict(orient="records") == [{"region": "South", "sales": 12}]
    assert frame.attrs["transformations"][0]["sheet"] == "Actual sales"


def test_parquet_reader_preserves_flat_records_and_format_metadata(tmp_path):
    path = tmp_path / "sales.parquet"
    parquet.write_table(pa.Table.from_pandas(pd.DataFrame({"region": ["South", "East"], "sales": [12, 15]})), path)

    frame = parse_file(path, "parquet")

    assert frame.to_dict(orient="records") == [
        {"region": "South", "sales": 12}, {"region": "East", "sales": 15}
    ]
    assert frame.attrs["source_format"] == "parquet"


def test_workbook_sheet_inspection_returns_nonempty_sheets(analysis_context):
    client, sessions, _ = analysis_context
    workbook = Workbook()
    workbook.active.title = "说明"
    workbook.active.append(["readme"])
    sheet = workbook.create_sheet("销售数据")
    sheet.append(["region", "sales"])
    sheet.append(["South", 12])
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()

    response = client.post("/api/v1/datasets/inspect-sheets", files={
        "file": ("sales.xlsx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})

    assert response.status_code == 200, response.text
    assert response.json()["data"] == {"sheets": ["销售数据"], "default": "销售数据"}
    upload = client.post("/api/v1/datasets/upload", data={"sheet_name": "销售数据"}, files={
        "file": ("sales.xlsx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert upload.status_code == 202, upload.text
    with sessions() as db:
        dataset = db.get(Dataset, upload.json()["data"]["id"])
        assert dataset.parse_options_json == {"sheet_name": "销售数据"}
