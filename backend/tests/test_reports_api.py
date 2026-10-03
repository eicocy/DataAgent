from datetime import UTC, datetime
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import select

from test_analysis_api import analysis_context, make_session
from app.models import AnalysisRecord, BackgroundJob, DatasetVersion, User
from app.reports.service import execute_report_task
from app.routers import analysis
from app.security import create_access_token, hash_password
from app.services.analysis_agent import AgentOutcome
from app.services.artifacts import ArtifactStore


def run_report_task(sessions, record_id):
    with sessions() as db:
        job = db.scalar(select(BackgroundJob).where(
            BackgroundJob.kind == "report", BackgroundJob.resource_id == record_id))
        job.status, job.lease_token = "running", "test-report-lease"
        db.commit()
        execute_report_task(db, record_id, job.id, job.lease_token)

class ReportSourceAgent:
    def analyze(self, *_):
        return AgentOutcome(
            tool_calls=[{"tool_call_id": "report-source-tool", "tool_name": "group_by_analysis",
                         "schema_version": "1.0", "parameters": {"group_columns": ["region"],
                         "value_column": "sales", "aggregation": "sum"}, "status": "succeeded",
                         "duration_ms": 2, "result_summary": "已汇总地区销售额"}],
            tool_result={"columns": ["region", "sales_sum"], "rows": [{"region": "East", "sales_sum": 12}],
                         "row_count": 1, "truncated": False},
            answer="East 销售额为 12。", chart=None, status="succeeded")


def test_report_is_versioned_exported_and_downloaded_with_owner_checks(analysis_context, monkeypatch, tmp_path):
    client, sessions, _ = analysis_context
    uploaded = client.post("/api/v1/datasets/upload", files={
        "file": ("sales.csv", b"region,sales\nEast,12\n", "text/csv")})
    dataset_id = uploaded.json()["data"]["id"]
    session_id = make_session(sessions, dataset_id)
    monkeypatch.setattr(analysis, "agent", ReportSourceAgent())
    run = client.post("/api/v1/analysis/chat", json={"session_id": session_id,
        "dataset_id": dataset_id, "question": "按地区汇总销售额", "request_id": "report-source-1"})
    assert run.status_code == 200
    record_id = run.json()["data"]["record_id"]
    monkeypatch.setattr("app.config.get_settings", lambda: type("Settings", (), {
        "artifact_dir": str(tmp_path / "artifacts"), "artifact_retention_days": 7})())
    with sessions() as db:
        record = db.get(AnalysisRecord, record_id)
        version = db.scalar(select(DatasetVersion).where(DatasetVersion.dataset_id == dataset_id))
        user = db.scalar(select(User).where(User.username == "alice"))
        full_result = {"columns": ["region", "sales_sum"], "rows": [
            {"region": "East", "sales_sum": 12}, {"region": "West", "sales_sum": 8}],
            "row_count": 2, "total": 2}
        stored_result = ArtifactStore(db).write(record, "complete_result", full_result)
        record.tool_result_json = full_result
        record.plan_json = {"version": "2.0", "steps": [{"step_id": "complete_result",
            "tool_name": "group_by_analysis", "status": "COMPLETED",
            "result_ref": f"artifact:{stored_result['artifact_id']}"}]}
        db.commit()
        report_spec = {"title": "销售分析报告", "report_type": "sales", "dataset_id": dataset_id,
                       "dataset_version_id": version.id}
        user_id = user.id
    cancelled_task = client.post("/api/v1/reports", json={"session_id": session_id,
        "source_record_ids": [record_id], "spec": {**report_spec, "title": "取消的报告"}})
    assert cancelled_task.status_code == 202
    cancelled_id = cancelled_task.json()["data"]["record_id"]
    cancelled = client.post(f"/api/v1/analysis/runs/{cancelled_id}/cancel")
    assert cancelled.status_code == 200
    assert client.get(f"/api/v1/analysis/runs/{cancelled_id}").json()["data"]["status"] == "cancelled"

    created = client.post("/api/v1/reports", json={"session_id": session_id,
        "source_record_ids": [record_id], "spec": report_spec})
    assert created.status_code == 202, created.text
    run_report_task(sessions, created.json()["data"]["record_id"])
    report_run = client.get(created.json()["data"]["status_url"])
    report = report_run.json()["data"]["report"]["report"]
    assert report["version"] == 1
    assert report["document"]["metadata"]["dataset_version_id"] == version.id

    changed = {**report_spec, "report_id": report["id"], "base_version": 1, "title": "销售分析报告修订"}
    updated = client.patch(f"/api/v1/reports/{report['id']}", json={"session_id": session_id,
        "source_record_ids": [record_id], "spec": changed})
    assert updated.status_code == 200
    assert updated.json()["data"]["version"] == 2
    conflict = client.patch(f"/api/v1/reports/{report['id']}", json={"session_id": session_id,
        "source_record_ids": [record_id], "spec": changed})
    assert conflict.status_code == 409

    exported = client.post(f"/api/v1/reports/{report['id']}/versions/1/exports/pdf")
    assert exported.status_code == 202, exported.text
    run_report_task(sessions, exported.json()["data"]["record_id"])
    export_run = client.get(exported.json()["data"]["status_url"])
    artifact_id = export_run.json()["data"]["report"]["artifacts"][0]["artifact_id"]
    cached = client.post(f"/api/v1/reports/{report['id']}/versions/1/exports/pdf")
    assert cached.status_code == 202
    run_report_task(sessions, cached.json()["data"]["record_id"])
    cached_run = client.get(cached.json()["data"]["status_url"])
    assert cached_run.json()["data"]["report"]["artifacts"][0]["artifact_id"] == artifact_id
    downloaded = client.get(f"/api/v1/artifacts/{artifact_id}/download")
    assert downloaded.status_code == 200
    assert downloaded.content.startswith(b"%PDF-")
    assert downloaded.headers["content-type"] == "application/pdf"
    ranged = client.get(f"/api/v1/artifacts/{artifact_id}/download", headers={"Range": "bytes=0-4"})
    assert ranged.status_code == 206
    assert ranged.content == b"%PDF-"

    xlsx_task = client.post(f"/api/v1/reports/{report['id']}/versions/1/exports/xlsx")
    assert xlsx_task.status_code == 202
    run_report_task(sessions, xlsx_task.json()["data"]["record_id"])
    xlsx_run = client.get(xlsx_task.json()["data"]["status_url"])
    xlsx_artifact_id = xlsx_run.json()["data"]["report"]["artifacts"][0]["artifact_id"]
    xlsx_download = client.get(f"/api/v1/artifacts/{xlsx_artifact_id}/download")
    workbook = load_workbook(BytesIO(xlsx_download.content), read_only=True, data_only=True)
    assert workbook["Raw Data"].max_row == 3
    data_row = list(workbook["Raw Data"].values)[1]
    assert data_row[:2] == ("East", 12)
    assert data_row[-1] == record_id
    workbook.close()

    client.post("/api/v1/auth/logout")
    with sessions() as db:
        now = datetime.now(UTC)
        other = User(username="bob", password_hash=hash_password("safe-password-456"),
                     status="active", created_at=now, updated_at=now)
        db.add(other)
        db.commit()
        other_id = other.id
    client.cookies.clear()
    client.cookies.set("datalens_session", create_access_token(other_id), path="/api/v1")
    assert client.get(f"/api/v1/reports/{report['id']}").status_code == 404
    assert client.get(f"/api/v1/artifacts/{artifact_id}").status_code == 404
    assert user_id > 0
