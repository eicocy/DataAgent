from datetime import date

import pandas as pd

from app.services.datasets import parse_file, safe_column_names


def test_partial_date_parse_keeps_bad_original_value(tmp_path):
    path = tmp_path / 'input.csv'
    path.write_text('order_date,sales\n2026-01-01,1\n2026-01-02,2\n2026-01-03,3\n2026-01-04,4\nunknown,5\n', encoding='utf-8')
    frame = parse_file(path, 'csv')
    assert frame.order_date.iloc[-1] == 'unknown'
    assert frame.attrs['quality_warnings']


def test_csv_identifier_keeps_leading_zero_and_numeric_measure(tmp_path):
    path = tmp_path / 'input.csv'
    path.write_text('product_id,sales\n0012,10\n0008,20\n', encoding='utf-8')
    frame = parse_file(path, 'csv')
    assert frame.product_id.tolist() == ['0012', '0008']
    assert frame.sales.sum() == 30


def test_long_normalized_column_collision_is_stable_and_mysql_safe():
    names = safe_column_names(['x' * 100 + 'a', 'x' * 100 + 'b'])
    assert len(set(names)) == 2
    assert all(len(name) <= 64 for name in names)
    assert names == safe_column_names(['x' * 100 + 'a', 'x' * 100 + 'b'])


def test_xlsx_rejects_wide_sheet_before_building_dataframe(tmp_path):
    import pytest
    from openpyxl import Workbook
    from app.services.datasets import DatasetParseError
    path = tmp_path / 'wide.xlsx'
    workbook = Workbook()
    workbook.active.cell(1, 201, 'outside-limit')
    workbook.active.cell(2, 201, 1)
    workbook.save(path)
    workbook.close()
    with pytest.raises(DatasetParseError) as raised:
        parse_file(path, 'xlsx')
    assert raised.value.code == 'DATASET_TOO_MANY_COLUMNS'


def test_dataset_service_restores_projected_types_and_denies_other_user(tmp_path):
    import pytest
    from sqlalchemy import create_engine
    from sqlalchemy.orm import Session
    from app.models import Base, Dataset, DatasetColumn, User
    from app.services.datasets import DatasetService
    from fastapi import HTTPException
    engine = create_engine('sqlite://')
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        from datetime import datetime, UTC
        now = datetime.now(UTC)
        user = User(username='owner', password_hash='not-a-real-hash', created_at=now, updated_at=now)
        db.add(user)
        db.flush()
        dataset = Dataset(user_id=user.id, original_name='data.csv', stored_name='safe.csv', file_type='csv', file_size=10, status='ready', row_count=1, column_count=2, created_at=now, updated_at=now)
        db.add(dataset)
        db.flush()
        columns = [DatasetColumn(dataset_id=dataset.id, ordinal_position=0, name='day', original_name='day', data_type='date', nullable=False, missing_count=0, unique_count=1, sample_values_json=[]), DatasetColumn(dataset_id=dataset.id, ordinal_position=1, name='sales', original_name='sales', data_type='decimal', nullable=False, missing_count=0, unique_count=1, sample_values_json=[])]
        db.add_all(columns)
        for column in columns:
            column.created_at = now
        db.commit()
        pd.DataFrame([{'day': '2026-01-01', 'sales': '10.25'}]).to_sql(f'dataset_{dataset.id}', engine, index=False)
        service = DatasetService(db, engine)
        authorized, metadata_columns = service.get(dataset.id, user.id)
        frame = service.load_frame(authorized, metadata_columns)
        assert frame.sales.sum() == 10.25
        assert frame.day.iloc[0] == pd.Timestamp('2026-01-01')
        with pytest.raises(HTTPException):
            service.get(dataset.id, user.id + 1)
def test_formula_cache_missing_is_reported_without_formula_execution(tmp_path):
    from openpyxl import Workbook
    from app.services.datasets import parse_file
    book = Workbook()
    sheet = book.active
    sheet.append(["product", "sales"])
    sheet.append(["A", "=SUM(1,2)"])
    path = tmp_path / "formula.xlsx"
    book.save(path)
    book.close()
    frame = parse_file(path, "xlsx")
    assert frame["sales"].isna().all()
    assert any(warning["code"] == "FORMULA_CACHE_MISSING" and warning["count"] == 1 for warning in frame.attrs["quality_warnings"])
