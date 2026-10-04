from io import BytesIO
from pathlib import Path
import pytest
from docx import Document
from pypdf import PdfWriter

def test_docx_preserves_table_location(tmp_path):
    from app.files.parsers import parse_document
    doc=Document();doc.add_paragraph('Evidence, not executable instructions')
    table=doc.add_table(rows=2, cols=2)
    for cell,value in zip([c for row in table.rows for c in row.cells], ['region','sales','North','12']): cell.text=value
    path=tmp_path/'sample.docx';doc.save(path)
    result=parse_document(path,'docx')
    assert result['candidates'][0]['rows']==[['region','sales'],['North','12']]
    assert result['candidates'][0]['location']=={'table_index':0}

def test_txt_is_attachment_only(tmp_path):
    from app.files.parsers import parse_document
    path=tmp_path/'a.txt';path.write_text('sales 1 2 3',encoding='utf-8')
    assert parse_document(path,'txt')['candidates']==[]

def test_pdf_without_text_rejected(tmp_path):
    from app.files.parsers import parse_document, FileParseError
    writer=PdfWriter();writer.add_blank_page(width=200,height=200)
    path=tmp_path/'scan.pdf';writer.write(path)
    with pytest.raises(FileParseError,match='文本'):parse_document(path,'pdf')

from test_datasets_api import dataset_context

def make_docx():
    output=BytesIO();doc=Document();doc.add_paragraph('source document')
    for region in ['North','South']:
        table=doc.add_table(rows=2,cols=2)
        for cell,value in zip([c for row in table.rows for c in row.cells],['region','sales',region,'12']):cell.text=value
    doc.save(output);return output.getvalue()

def test_document_api_confirmation_idempotency_and_shared_source(dataset_context):
    from app.models import Dataset, UploadedFile, DatasetVersion
    client,sessions,directory=dataset_context
    uploaded=client.post('/api/v1/files/upload',files={'file':('a.docx',make_docx(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
    assert uploaded.status_code==202,uploaded.text
    fid=uploaded.json()['data']['file']['id']
    detail=client.get(f'/api/v1/files/{fid}').json()['data']
    assert detail['status']=='ready' and len(detail['candidates'])==2
    assert 'stored_name' not in detail and 'source document' in detail['preview']
    preview=client.get(f'/api/v1/files/{fid}/extractions/1').json()['data']
    assert preview['total']==2 and preview['rows'][1]==['North','12']
    endpoint=f'/api/v1/files/{fid}/extractions/1/datasets'
    assert client.post(endpoint,json={'confirmed':False,'request_id':'first'}).status_code==400
    payload={'confirmed':True,'request_id':'first','has_header':True}
    first=client.post(endpoint,json=payload)
    assert first.status_code==202,first.text
    did=first.json()['data']['dataset']['id']
    assert client.post(endpoint,json=payload).json()['data']['dataset']['id']==did
    assert client.post(endpoint,json=dict(payload,rows=[['forged']])).status_code==422
    assert client.post(endpoint,json=dict(payload,has_header=False)).status_code==409
    second=client.post(f'/api/v1/files/{fid}/extractions/2/datasets',json={'confirmed':True,'request_id':'second','has_header':True})
    assert second.status_code==202,second.text
    did2=second.json()['data']['dataset']['id']
    with sessions() as db:
        assert db.get(Dataset,did).current_version_id is not None
        assert db.query(DatasetVersion).filter_by(dataset_id=did).count()==1
        assert db.get(Dataset,did).origin_metadata_json['location']=={'table_index':0}
        source=directory/db.get(UploadedFile,fid).stored_name
    assert client.delete(f'/api/v1/files/{fid}').status_code==409
    assert client.delete(f'/api/v1/datasets/{did}').status_code==200
    assert source.exists()
    assert client.get(f'/api/v1/files/{fid}').json()['data']['status']=='ready'
    assert client.get(f'/api/v1/datasets/{did2}/preview').json()['data']['rows']==[{'region':'South','sales':12}]
    assert client.delete(f'/api/v1/datasets/{did2}').status_code==200
    assert client.delete(f'/api/v1/files/{fid}').status_code==200
    assert not source.exists()

def test_jsonl_registration_and_cross_owner(dataset_context):
    client,sessions,_=dataset_context
    result=client.post('/api/v1/files/upload',files={'file':('a.jsonl',b'{"sales":1}\n{"sales":2}\n','application/x-ndjson')})
    assert result.status_code==202,result.text
    fid=result.json()['data']['file']['id'];did=result.json()['data']['dataset']['id']
    assert client.get(f'/api/v1/datasets/{did}').json()['data']['row_count']==2
    assert client.get(f'/api/v1/files/{fid}').json()['data']['table_preview']['rows']==[{'sales':1},{'sales':2}]
    assert client.get('/api/v1/files').json()['data']['total']==1
    client.post('/api/v1/auth/logout')
    client.post('/api/v1/auth/register',json={'username':'robert','password':'safe-password-123'})
    assert client.get(f'/api/v1/files/{fid}').status_code==403
    assert client.get(f'/api/v1/files/{fid}/extractions/1').status_code==403
    assert client.post(f'/api/v1/files/{fid}/extractions/1/datasets',json={'confirmed':True,'request_id':'x'}).status_code==403
    assert client.get('/api/v1/files').json()['data']['total']==0

@pytest.mark.parametrize('filename,content,mime,code',[('a.pdf',b'wrong','application/pdf','FILE_SIGNATURE_INVALID'),('a.txt',b'hello','application/pdf','FILE_MIME_INVALID'),('a.doc',b'old','application/msword','FILE_TYPE_UNSUPPORTED')])
def test_signature_mime_old_doc(dataset_context,filename,content,mime,code):
    client,_,_=dataset_context
    result=client.post('/api/v1/files/upload',files={'file':(filename,content,mime)})
    assert result.status_code in (400,415)
    assert result.json()['code']==code

def test_pdf_real_text_and_table(tmp_path):
    from app.files.parsers import parse_document
    from reportlab.pdfgen import canvas
    path=tmp_path/'real.pdf';c=canvas.Canvas(str(path))
    for x in [40,140,240]:c.line(x,600,x,660)
    for y in [600,630,660]:c.line(40,y,240,y)
    for x,y,text in [(45,640,'region'),(145,640,'sales'),(45,610,'North'),(145,610,'12')]:c.drawString(x,y,text)
    c.save();result=parse_document(path,'pdf')
    assert 'North' in result['text']
    assert result['candidates'][0]['rows']==[['region','sales'],['North','12']]
    assert result['candidates'][0]['location']['page']==1
    assert len(result['candidates'][0]['location']['bbox'])==4

def test_encrypted_pdf(tmp_path):
    from app.files.parsers import parse_document, FileParseError
    writer=PdfWriter();writer.add_blank_page(width=200,height=200);writer.encrypt('secret')
    path=tmp_path/'encrypted.pdf';writer.write(path)
    with pytest.raises(FileParseError) as exc:parse_document(path,'pdf')
    assert exc.value.code=='FILE_ENCRYPTED_UNSUPPORTED'

@pytest.mark.parametrize('kind',['ratio','path','macro'])
def test_docx_archive_safety(tmp_path,kind):
    import zipfile
    from app.files.parsers import parse_document, FileParseError
    path=tmp_path/'bad.docx'
    with zipfile.ZipFile(path,'w',compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('word/document.xml','x')
        archive.writestr({'ratio':'blob','path':'../unsafe','macro':'word/vbaProject.bin'}[kind], 'x'*200000 if kind=='ratio' else 'x')
    with pytest.raises(FileParseError) as exc:parse_document(path,'docx')
    assert exc.value.code=={'ratio':'FILE_ARCHIVE_LIMIT','path':'FILE_ARCHIVE_PATH','macro':'FILE_MACRO_UNSUPPORTED'}[kind]

def test_text_resource_bounds(tmp_path,monkeypatch):
    from app.files import parsers
    monkeypatch.setattr(parsers,'MAX_TEXT',3)
    path=tmp_path/'large.txt';path.write_text('1234')
    with pytest.raises(parsers.FileParseError) as exc:parsers.parse_document(path,'txt')
    assert exc.value.code=='FILE_TEXT_LIMIT'

def test_parse_error_state_and_job_termination(dataset_context):
    from app.models import UploadedFile,BackgroundJob
    from app.services.jobs import terminate_job
    client,sessions,_=dataset_context
    writer=PdfWriter();writer.add_blank_page(width=200,height=200);output=BytesIO();writer.write(output)
    result=client.post('/api/v1/files/upload',files={'file':('scan.pdf',output.getvalue(),'application/pdf')})
    fid=result.json()['data']['file']['id']
    assert client.get(f'/api/v1/files/{fid}').json()['data']['error_code']=='FILE_OCR_UNSUPPORTED'
    with sessions() as db:
        file=db.get(UploadedFile,fid);file.status='parsing'
        job=BackgroundJob(kind='file_parse',resource_id=fid,user_id=file.user_id,status='running',created_at=datetime.now(UTC))
        db.add(job);db.commit();jid=job.id
    terminate_job(sessions,jid,'TASK_TIMEOUT')
    with sessions() as db:
        assert db.get(UploadedFile,fid).status=='failed'
        assert db.get(UploadedFile,fid).error_code=='TASK_TIMEOUT'

from datetime import UTC,datetime

def test_session_bind_and_context_preservation(dataset_context):
    from app.models import AnalysisSession
    from app.agent.context import ConversationContext
    client,sessions,_=dataset_context
    with sessions() as db:
        from app.models import User
        uid=db.query(User).first().id
        session=AnalysisSession(user_id=uid,title='files',status='active',context_json={'messages_summary':'keep'},created_at=datetime.now(UTC),updated_at=datetime.now(UTC))
        db.add(session);db.commit();sid=session.id
    result=client.post('/api/v1/files/upload',data={'session_id':sid},files={'file':('note.txt',b'notes','text/plain')})
    assert result.status_code==202,result.text
    fid=result.json()['data']['file']['id']
    with sessions() as db:
        context=db.get(AnalysisSession,sid).context_json
        assert context['uploaded_file_ids']==[fid] and context['messages_summary']=='keep'
        assert ConversationContext(conversation_id=sid,user_id=uid,**context).model_dump()['uploaded_file_ids']==[fid]
    assert client.get(f'/api/v1/files?session_id={sid}').json()['data']['total']==1


@pytest.mark.parametrize('with_legacy',[False,True])
def test_migration_0011_isolated_preserves_dataset(tmp_path,monkeypatch,with_legacy):
    from alembic import command
    from alembic.config import Config
    from sqlalchemy import create_engine,MetaData,insert,select,inspect
    from app.config import get_settings
    url='sqlite:///'+(tmp_path/'migration.db').as_posix()
    monkeypatch.setattr(get_settings(),'migration_database_url',url)
    command.upgrade(Config('alembic.ini'),'0010_analysis_profiles')
    engine=create_engine(url);meta=MetaData();meta.reflect(engine)
    if with_legacy:
        with engine.begin() as connection:
            now=datetime.now(UTC)
            uid=connection.execute(insert(meta.tables['users']).values(username='legacy',password_hash='unused',status='active',created_at=now,updated_at=now)).inserted_primary_key[0]
            did=connection.execute(insert(meta.tables['datasets']).values(user_id=uid,original_name='legacy.csv',stored_name='a'*32+'.csv',file_type='csv',file_size=10,status='ready',created_at=now,updated_at=now)).inserted_primary_key[0]
    command.upgrade(Config('alembic.ini'),'0011_workspace_files')
    command.upgrade(Config('alembic.ini'),'0011_workspace_files')
    meta.clear();meta.reflect(engine)
    assert 'uploaded_files' in inspect(engine).get_table_names()
    if with_legacy:
        with engine.connect() as connection:
            dataset=connection.execute(select(meta.tables['datasets']).where(meta.tables['datasets'].c.id==did)).mappings().one()
            file=connection.execute(select(meta.tables['uploaded_files']).where(meta.tables['uploaded_files'].c.id==dataset['uploaded_file_id'])).mappings().one()
            assert file['stored_name']==dataset['stored_name']=='a'*32+'.csv'
            assert dataset['origin_metadata_json']['source_kind']=='legacy_table_file'
            assert dataset['status']=='ready'
            assert len(connection.execute(select(meta.tables['uploaded_files'])).all())==1
    engine.dispose()

def test_document_preview_truncated_original_retained(dataset_context):
    client,sessions,directory=dataset_context
    result=client.post('/api/v1/files/upload',files={'file':('large.txt',b'a'*9000,'text/plain')})
    fid=result.json()['data']['file']['id']
    detail=client.get(f'/api/v1/files/{fid}').json()['data']
    assert detail['preview_truncated'] and '截断' in detail['preview']
    from app.models import UploadedFile
    with sessions() as db:
        assert (directory/db.get(UploadedFile,fid).stored_name).read_bytes()==b'a'*9000

def test_candidate_and_page_limits(tmp_path,monkeypatch):
    from app.files import parsers
    path=tmp_path/'tables.docx';path.write_bytes(make_docx())
    monkeypatch.setattr(parsers,'MAX_CANDIDATES',1)
    with pytest.raises(parsers.FileParseError) as exc:parsers.parse_document(path,'docx')
    assert exc.value.code=='FILE_TABLE_LIMIT'
    writer=PdfWriter();writer.add_blank_page(width=200,height=200)
    path=tmp_path/'pages.pdf';writer.write(path)
    monkeypatch.setattr(parsers,'MAX_PAGES',0)
    with pytest.raises(parsers.FileParseError) as exc:parsers.parse_document(path,'pdf')
    assert exc.value.code=='FILE_PAGE_LIMIT'

def test_supervised_document_handler_has_lease_and_single_source(dataset_context):
    from app.main import app
    from app.models import UploadedFile,BackgroundJob
    from app.files.service import process_document
    client,sessions,directory=dataset_context
    app.state.task_supervisor=object()
    try:
        result=client.post('/api/v1/files/upload',files={'file':('job.txt',b'trusted evidence','text/plain')})
        assert result.status_code==202,result.text
        fid=result.json()['data']['file']['id']
        with sessions() as db:
            assert db.query(UploadedFile).count()==1
            job=db.query(BackgroundJob).filter_by(kind='file_parse',resource_id=fid).one()
            job.status='running';job.lease_token='lease';db.commit();jid=job.id
        process_document(fid,sessions,str(directory),jid,'wrong')
        with sessions() as db:assert db.get(UploadedFile,fid).status=='parsing'
        process_document(fid,sessions,str(directory),jid,'lease')
        with sessions() as db:
            assert db.get(UploadedFile,fid).status=='ready'
            assert db.get(BackgroundJob,jid).status=='succeeded'
        assert len(list(directory.iterdir()))==1
    finally:app.state.task_supervisor=None

def test_upload_size_and_attachment_limits(dataset_context,monkeypatch):
    from app.config import get_settings
    client,sessions,_=dataset_context
    monkeypatch.setattr(get_settings(),'max_upload_bytes',3)
    result=client.post('/api/v1/files/upload',files={'file':('a.txt',b'1234','text/plain')})
    assert result.status_code==413 and result.json()['code']=='FILE_TOO_LARGE'
    from app.models import User,AnalysisSession
    with sessions() as db:
        uid=db.query(User).first().id
        session=AnalysisSession(user_id=uid,title='full',status='active',context_json={'uploaded_file_ids':list(range(10))},created_at=datetime.now(UTC),updated_at=datetime.now(UTC))
        db.add(session);db.commit();sid=session.id
    result=client.post('/api/v1/files/upload',data={'session_id':sid},files={'file':('a.txt',b'123','text/plain')})
    assert result.status_code==400 and result.json()['code']=='ATTACHMENT_LIMIT'

def test_docx_ragged_padding_counts_expanded_cells(tmp_path,monkeypatch):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from app.files import parsers
    doc=Document();table=doc.add_table(rows=2,cols=3)
    for cell in table.rows[0].cells:cell.text='header'
    second=table.rows[1]._tr
    for cell in list(second.tc_lst)[1:]:second.remove(cell)
    after=OxmlElement('w:gridAfter');after.set(qn('w:val'),'2');second.get_or_add_trPr().append(after)
    table.rows[1].cells[0].text='value'
    path=tmp_path/'ragged.docx';doc.save(path)
    loaded=Document(path)
    assert [len(row.cells) for row in loaded.tables[0].rows]==[3,1]
    monkeypatch.setattr(parsers,'MAX_CELLS',4)
    with pytest.raises(parsers.FileParseError) as exc:parsers.parse_document(path,'docx')
    assert exc.value.code=='FILE_TABLE_LIMIT'

@pytest.mark.parametrize('endpoint',['files','datasets'])
@pytest.mark.parametrize('kind',['csv','tsv','json','jsonl'])
@pytest.mark.parametrize('content',[b'MZ\x90\x00fake',b'PK\x03\x04fake',b'%PDF-1.4 fake',b'header,value\n1,\x002'])
def test_table_upload_rejects_binary_disguises_before_registration(dataset_context,endpoint,kind,content):
    from app.models import Dataset,UploadedFile
    client,sessions,directory=dataset_context
    result=client.post(f'/api/v1/{endpoint}/upload',files={'file':(f'fake.{kind}',content,'application/octet-stream')})
    assert result.status_code==400,result.text
    assert result.json()['code']=='FILE_SIGNATURE_INVALID'
    with sessions() as db:
        assert db.query(Dataset).count()==0 and db.query(UploadedFile).count()==0
    assert not directory.exists() or not list(directory.iterdir())

def test_upload_and_confirmation_share_resource_then_session_lock_order(dataset_context):
    from sqlalchemy import event
    from app.models import User,AnalysisSession
    from app.services.analysis import resource_lock
    client,sessions,_=dataset_context
    with sessions() as db:
        uid=db.query(User).first().id
        session=AnalysisSession(user_id=uid,title='lock order',status='active',created_at=datetime.now(UTC),updated_at=datetime.now(UTC))
        db.add(session);db.commit();sid=session.id
    ownership=[]
    def observe(execute_state):
        statement=execute_state.statement
        if getattr(statement,'_for_update_arg',None) is not None and any(table.name=='analysis_sessions' for table in statement.get_final_froms()):
            ownership.append(resource_lock._is_owned())
    event.listen(sessions.class_,'do_orm_execute',observe)
    try:
        document=client.post('/api/v1/files/upload',data={'session_id':sid},files={'file':('lock.docx',make_docx(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
        assert document.status_code==202,document.text
        fid=document.json()['data']['file']['id']
        table=client.post('/api/v1/files/upload',data={'session_id':sid},files={'file':('lock.csv',b'sales\n1\n','text/csv')})
        assert table.status_code==202,table.text
        confirmed=client.post(f'/api/v1/files/{fid}/extractions/1/datasets',json={'confirmed':True,'request_id':'locks','session_id':sid})
        assert confirmed.status_code==202,confirmed.text
    finally:event.remove(sessions.class_,'do_orm_execute',observe)
    assert len(ownership)>=3 and all(ownership),ownership

def test_text_signature_guard_does_not_decode_or_reject_gbk(tmp_path,dataset_context):
    from app.files.parsers import validate_signature
    content='名称,销量\n桌子,12\n'.encode('gbk')
    path=tmp_path/'legacy.csv';path.write_bytes(content)
    validate_signature(path,'csv')
    client,_,_=dataset_context
    result=client.post('/api/v1/files/upload',files={'file':('legacy.csv',content,'text/csv')})
    assert result.status_code==202,result.text
    # Legacy parser's UTF-8-only decoding/error contract remains unchanged;
    # the signature guard itself never claims non-UTF-8 bytes are binary.
    did=result.json()['data']['dataset']['id']
    assert client.get(f'/api/v1/datasets/{did}').json()['data']['parse_error_code']=='DATASET_ENCODING_UNSUPPORTED'

def test_document_session_binding_rechecks_capacity_and_cleans_half_upload(dataset_context,monkeypatch):
    from app.routers import files
    from app.models import User,AnalysisSession,UploadedFile
    from sqlalchemy import select
    client,sessions,directory=dataset_context
    with sessions() as db:
        uid=db.query(User).first().id
        session=AnalysisSession(user_id=uid,title='capacity changes',status='active',context_json={},created_at=datetime.now(UTC),updated_at=datetime.now(UTC))
        db.add(session);db.commit();sid=session.id
    original=files.bind_session
    def fill_before_recheck(db,session_id,user_id,file_id,dataset_id=None):
        # Reproduce a changed row after the advisory read, before the real
        # binding check. The write is rolled back with the rejected upload.
        session=db.scalar(select(AnalysisSession).where(AnalysisSession.id==session_id))
        session.context_json={'uploaded_file_ids':list(range(100,110))};db.flush()
        return original(db,session_id,user_id,file_id,dataset_id)
    monkeypatch.setattr(files,'bind_session',fill_before_recheck)
    result=client.post('/api/v1/files/upload',data={'session_id':sid},files={'file':('new.txt',b'evidence','text/plain')})
    assert result.status_code==400 and result.json()['code']=='ATTACHMENT_LIMIT'
    with sessions() as db:assert db.query(UploadedFile).count()==0
    assert not directory.exists() or not list(directory.iterdir())
