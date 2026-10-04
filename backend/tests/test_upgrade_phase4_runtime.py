from sqlalchemy import select
from test_analysis_api import analysis_context
from app.models import AnalysisRecord, Dataset, BackgroundJob
from app.services.analysis import execute_record
from app.services.analysis_agent import DeepSeekAgent
from app.agent.providers import FakeLLMProvider


def test_real_analysis_pipeline_delivers_and_does_not_complete_early(analysis_context,monkeypatch):
    client,sessions,_ = analysis_context
    did = client.post('/api/v1/datasets/upload',files={'file':('a.csv',b'x\n1\n2\n','text/csv')}).json()['data']['id']
    sid = client.post('/api/v1/analysis/sessions',json={'dataset_id':did}).json()['data']['id']
    accepted = client.post('/api/v1/analysis/runs',json={'session_id':sid,'dataset_id':did,
        'question':'检查质量并生成报告','request_id':'phase4-runtime','depth':'STANDARD',
        'profile_ids':['general-quality'],'report_template':'data_quality','output_formats':['xlsx','docx','pdf','python']})
    assert accepted.status_code==202,accepted.text
    rid = accepted.json()['data']['record_id']
    with sessions() as db:
        vid=db.get(Dataset,did).current_version_id
        provider=FakeLLMProvider([
            {'intent':'DATA_ANALYSIS','confidence':.99,'requires_dataset':True,'requires_analysis':True},
            {'task_id':str(rid),'goal':'quality','intent':'DATA_ANALYSIS','dataset_id':did,'dataset_version_id':vid,
             'inputs':[{'alias':'primary','dataset_id':did,'dataset_version_id':vid}],
             'steps':[{'step_id':'quality','tool_name':'missing_value_analysis'}]},
            {'template':'字段缺失数为 {count}。','facts':[{'key':'count','step_id':'quality','path':['findings',0,'count']}],'evidence_refs':['quality']}])
        from app.reports import delivery
        original = delivery.export_report_version
        def observed(db,report,version,format):
            assert db.get(AnalysisRecord,rid).status=='running'
            return original(db,report,version,format)
        monkeypatch.setattr(delivery,'export_report_version',observed)
        execute_record(db,rid,DeepSeekAgent(provider_factory=lambda emit:provider),db.get_bind(),db.get_bind())
        row=db.get(AnalysisRecord,rid)
        assert row.status=='succeeded',(row.error_code,row.report_json)
        assert row.report_json['delivery']['status']=='succeeded'
        assert len(row.report_json['delivery']['artifacts'])==5
        assert not db.scalar(select(BackgroundJob).where(BackgroundJob.kind=='report'))
        assert row.report_json['delivery']['report']['document']['evidence'][0]['value']==0
        # One small end-to-end delivery smoke also opens the authenticated
        # previews and downloads, rather than repeating a separate export run.
        import io
        from openpyxl import load_workbook
        from docx import Document
        from pypdf import PdfReader
        for artifact in row.report_json['delivery']['artifacts']:
            response=client.get(artifact['download_url'])
            assert response.status_code==200
            preview=client.get(artifact['preview_url'])
            assert preview.status_code==200
            if artifact['type']=='excel': assert load_workbook(io.BytesIO(response.content))['Raw Data'].max_row==3
            elif artifact['type']=='word': assert Document(io.BytesIO(response.content)).paragraphs
            elif artifact['type']=='pdf': assert len(PdfReader(io.BytesIO(response.content)).pages)>=2
        from app.models import AnalysisArtifact, ToolExecutionRecord
        aid=db.scalar(select(AnalysisArtifact.id).where(AnalysisArtifact.record_id==rid,AnalysisArtifact.step_id=='quality'))
        execution_ids=db.scalars(select(ToolExecutionRecord.id)).all()
    followup=client.post('/api/v1/analysis/runs',json={'session_id':sid,'dataset_id':did,
        'question':'复用质量结果','request_id':'phase4-reference-runtime','depth':'STANDARD',
        'profile_ids':['general-quality'],'artifact_refs':[aid]})
    assert followup.status_code==202,followup.text
    next_id=followup.json()['data']['record_id']
    with sessions() as db:
        provider=FakeLLMProvider([
            {'intent':'DATA_ANALYSIS','confidence':.99,'requires_dataset':True,'requires_analysis':True},
            {'task_id':str(next_id),'goal':'quality','intent':'DATA_ANALYSIS','dataset_id':did,'dataset_version_id':vid,
             'inputs':[{'alias':'primary','dataset_id':did,'dataset_version_id':vid}],
             'steps':[{'step_id':'quality','tool_name':'missing_value_analysis'}]},
            {'template':'字段缺失数为 {count}。','facts':[{'key':'count','step_id':'quality','path':['findings',0,'count']}],'evidence_refs':['quality']}])
        execute_record(db,next_id,DeepSeekAgent(provider_factory=lambda emit:provider),db.get_bind(),db.get_bind())
        replay=db.get(AnalysisRecord,next_id)
        assert replay.status=='succeeded',(replay.error_code,replay.report_json)
        assert replay.tool_calls_json[0]['reused'] is True
        assert replay.plan_json['steps'][0]['result_ref']==f'artifact:{aid}'
        audit=db.scalars(select(ToolExecutionRecord).where(ToolExecutionRecord.id.not_in(execution_ids))).all()
        assert len(audit)==1 and audit[0].result_json['reused'] is True
        assert audit[0].result_json['artifact_ref']==aid
